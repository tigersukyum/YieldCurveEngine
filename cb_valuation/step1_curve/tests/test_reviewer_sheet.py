# -*- coding: utf-8 -*-
"""
검토자 수식 시트(io/reviewer_sheet) 테스트.
1. 파이썬 시트 모형(evaluate) vs fixture C(2024-12-31 검토자 시트 값복사 표본) — 원본은 값복사(해찾기 잔차)라 행별 허용오차를 둔다.
2. 수식 문자열 자체(cells) — 원본에 남은 수식 5개·앵커 체인·순환 참조 없음.
3. REVIEWER_2024 실행 → write_xlsx 가 Rf_dc/Rd_dc 를 수식 시트로 쓰고(수식 셀 확인), 값 시트 폴백(일간 격자)도 동작.
4. 서식 대응(_orig_col)이 열 성격을 보존한다.
Excel 실제 재계산 대조는 reference/verify_reviewer_sheet.py(Excel 필요)로 한다.
"""
import json, os, shutil, tempfile, unittest
from datetime import datetime

from cb_valuation.step1_curve.io import reviewer_sheet as RS
from cb_valuation.step1_curve.tests.test_app_pipeline import drive_to_done

HERE = os.path.dirname(os.path.abspath(__file__))
FIX = os.path.join(HERE, "fixtures", "reviewer_curves_20241231.json")
# 10000 배율 행·해찾기 잔차·미확인 원본 정의 행은 허용오차를 따로 둔다(원본 값복사와의 차이 근거: docs/REVIEWER_SHEET_SPEC.md §5)
TOL_RATE, TOL_FACE = 2e-13, 1e-8
SKIP_ROWS = {"RF": {20, 38}, "RD": {36, 52}}  # r20 AM 열 원본 이상치(31ulp), r38/r36 역방향 PV 출발값 미확인(원본 TC38=#REF!), r52 원본 정의 미확인
FACE_ROWS = {"RF": {21, 28}, "RD": {19, 26, 40, 46}}


def _col_index(s: str) -> int:
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def params_from_fixture(kind: str) -> RS.Params:
    with open(FIX, encoding="utf-8") as fh:
        d = json.load(fh)["sheets"][RS.SHEET_NAMES[kind]]
    steps = [int(d["row4"]["cells"][c]) for c in "CDEFGHIJKLMN"]
    ytm = [float(d["row6"]["cells"][c]) for c in "CDEFGHIJKLMN"]
    tenors = [{"label": RS._tenor_label(w / 52), "years": w / 52, "step": w, "ytm": y} for w, y in zip(steps, ytm)]
    return RS.Params(kind, 52, 4, 520, tenors, datetime(2024, 12, 31), header="fixture C")


class TestModelVsFixture(unittest.TestCase):
    def test_model_matches_reviewer_values(self):
        for kind in ("RF", "RD"):
            p = params_from_fixture(kind); V = RS.evaluate(p)
            with open(FIX, encoding="utf-8") as fh:
                d = json.load(fh)["sheets"][p.sheet]
            checked = 0
            for rk, row in d.items():
                r = int(rk[3:])
                if r in SKIP_ROWS[kind] or r not in V:
                    continue
                tol = TOL_FACE if r in FACE_ROWS[kind] else TOL_RATE
                for cc, v in row["cells"].items():
                    if not isinstance(v, (int, float)) or isinstance(v, bool):
                        continue
                    ci = _col_index(cc) - 3
                    mv = V[(r, "X")] if ci == p.N else (V[r][ci] if ci < len(V[r]) else None)
                    if mv is None or isinstance(mv, bool):
                        continue
                    if kind == "RD" and r == 50 and ci == p.N:
                        continue  # 원본 TC50 은 오래된 값(0.1549 ≠ PRODUCT) — 명세 §5
                    self.assertAlmostEqual(mv, v, delta=tol + 1e-15 * abs(v), msg=f"{p.sheet}!{cc}{r}: 모형 {mv!r} vs 원본 {v!r}")
                    checked += 1
            self.assertGreater(checked, 500, kind)

    def test_checks_true(self):
        for kind in ("RF", "RD"):
            p = params_from_fixture(kind); V = RS.evaluate(p); R = p.rows
            self.assertIs(V[R["mc1"]][0], True); self.assertIs(V[R["mc2"]][0], True)
            if kind == "RD":
                self.assertIs(V[R["wmc1"]][0], True); self.assertIs(V[R["wmc2"]][0], True)
            self.assertLess(max(abs(x - p.face) for x in V[R["qmc"]]), 1e-8)  # par 검증(MODEL CHECK = 액면)


class TestFormulas(unittest.TestCase):
    def test_formula_strings(self):
        p = params_from_fixture("RF"); c = RS.cells(p)
        self.assertEqual(c[(22, 3)], "=(C21-C19*C21*C23)/(C21+C19*C21)")
        self.assertEqual(c[(28, 3)], "=C19*C21*C23+(C19*C21+C21)*C22")
        self.assertEqual(c[(18, 3)], "=HLOOKUP(C17,$C$10:$TB$11,2,FALSE)")
        self.assertEqual(c[(13, 3)], "=((1+$O$32)^13)^(1/13)-1")
        self.assertEqual(c[(13, 16)], "=((1+$AB$32)^26/(1+$O$32)^13)^(1/13)-1")  # P13 (주 14, Q2)
        self.assertEqual(c[(32, 15)], "=(1/C22)^(1/O31)-1")  # 분기말 앵커
        self.assertEqual(c[(32, 4)], "=(D34*D33)^(1/D31)-1")
        self.assertEqual(c[(32, 523)], "=1/(1+TB32)^TB31"); self.assertEqual(c[(36, 523)], "=PRODUCT(C36:TB36)")
        self.assertEqual(c[(12, 3)], "=(1+C32)^52-1"); self.assertEqual(c[(7, 3)], "=INDEX($C$12:$TB$12,C$4)")
        self.assertEqual(c[(25, 3)], "=C18"); self.assertEqual(c[(25, 4)], "=(1+D24)^4-1")
        self.assertEqual(c[(21, 3)], 10000); self.assertEqual(c[(23, 3)], 0); self.assertEqual(c[(34, 3)], 1)
        self.assertTrue(c[(11, 3)].startswith("=IF(C$10<=$C$4,$C$6,IF(C$10>=$N$4,$N$6,INDEX($C$6:$N$6,MATCH(C$10,$C$4:$N$4,1))"))
        self.assertEqual(c[(9, 15)], "3mth"); self.assertEqual(c[(30, 15)], "0.25YR"); self.assertEqual(c[(30, 522)], "10year")
        self.assertEqual(c[(34, 2)], "FOMULA II -CUMM"); self.assertEqual(c[(32, 2)], "WEEKLY SPOT RATE ")  # 원본 철자·끝 공백 그대로
        q = RS.cells(params_from_fixture("RD"))
        self.assertEqual(q[(19, 3)], "=PV(C18/4,C16,-C18/4*$C$40,-$C$40)")
        self.assertEqual(q[(20, 3)], "=(C19-C18/4*C19*C21)/(C19+C18/4*C19)")
        self.assertEqual(q[(34, 3)], "=(1+C33)^52-1"); self.assertEqual(q[(43, 3)], "=($C$40-C42*$C$40*C44)/($C$40+C42*$C$40)")
        self.assertEqual(q[(47, 3)], "=(1+C45)^C41/C48"); self.assertEqual(q[(36, 522)], "=$C$40*TB37"); self.assertEqual(q[(40, 3)], 10000)

    def test_no_circular_reference(self):
        """수식 참조 그래프(정적)에 순환이 없다 — 분기말 앵커(r32)는 이표 격자 DF 만 참조한다."""
        import re
        for kind in ("RF", "RD"):
            p = params_from_fixture(kind); cs = RS.cells(p)
            ref = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
            graph = {}
            for (r, c), v in cs.items():
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                deps = set()
                for m in ref.finditer(v):
                    a, r1, b, r2 = m.group(1), int(m.group(2)), m.group(3), m.group(4)
                    if b:
                        for rr in range(r1, int(r2) + 1):
                            for cc in range(_col_index(a), _col_index(b) + 1):
                                deps.add((rr, cc))
                    else:
                        deps.add((r1, _col_index(a)))
                graph[(r, c)] = deps
            state = {}

            def visit(n, stack):
                if state.get(n) == 2:
                    return
                if state.get(n) == 1:
                    raise AssertionError(f"{kind}: 순환 참조 {stack[-3:]} → {n}")
                state[n] = 1
                for d in graph.get(n, ()):
                    if d in graph:
                        visit(d, stack + [n])
                state[n] = 2
            import sys
            sys.setrecursionlimit(10000)
            for n in graph:
                visit(n, [])


class TestStyleMapping(unittest.TestCase):
    def test_orig_col_classes(self):
        p = params_from_fixture("RF")
        self.assertEqual(RS._orig_col("weekly", 3, p), 3)      # 첫 주
        self.assertEqual(RS._orig_col("weekly", 15, p), 15)    # 주 13 = q1·테너·분기말
        self.assertEqual(RS._orig_col("weekly", 16, p), 16)    # 주 14 = 분기 첫 주
        self.assertEqual(RS._orig_col("weekly", 28, p), 28)    # 주 26 = 테너·분기말
        self.assertEqual(RS._orig_col("weekly", 67, p), 67)    # 주 65 = 분기말(테너 아님)
        self.assertEqual(RS._orig_col("weekly", 522, p), 522)  # 마지막
        self.assertEqual(RS._orig_col("weekly", 523, p), 523); self.assertIsNone(RS._orig_col("weekly", 600, p))
        self.assertEqual(RS._orig_col("quarterly", 3, p), 3); self.assertEqual(RS._orig_col("quarterly", 7, p), 7); self.assertEqual(RS._orig_col("quarterly", 42, p), 42)
        self.assertEqual(RS._orig_col("tenor", 3, p), 3); self.assertEqual(RS._orig_col("tenor", 14, p), 14); self.assertEqual(RS._orig_col("tenor", 15, p), 15)
        # 다른 격자(5년, 260주): 마지막 열은 원본 TB 서식, 분기말은 BO/AB 서식
        p5 = RS.Params("RF", 52, 4, 260, p.tenors[:10], None)
        self.assertEqual(RS._orig_col("weekly", 262, p5), 522); self.assertEqual(RS._orig_col("weekly", 67, p5), 67)


class TestWriteXlsx(unittest.TestCase):
    def test_reviewer_profile_writes_formula_sheets(self):
        from openpyxl import load_workbook
        from cb_valuation.step1_curve.io.evidence_writer import export_xlsx, dc_blocks
        base = tempfile.mkdtemp(prefix="cb_rs_")
        try:
            s, C = drive_to_done(base, "REVIEWER_2024", step="weekly", horizon_years=10)
            self.assertEqual(s.run.status, "done", s.result.fail_reason)
            self.assertTrue(RS.applicable(s, C)[0], RS.applicable(s, C)[1])
            path = export_xlsx(s, C, base)
            wb = load_workbook(path)
            ws = wb["Rf_dc"]
            self.assertEqual(ws["C22"].value, "=(C21-C19*C21*C23)/(C21+C19*C21)")
            self.assertEqual(ws["B32"].value, "WEEKLY SPOT RATE "); self.assertEqual(ws.freeze_panes, "E1"); self.assertFalse(ws.sheet_view.showGridLines)
            self.assertAlmostEqual(ws["C6"].value, s.grid.knots["RF"][0]["ytm"], delta=1e-15); self.assertEqual(ws["C4"].value, 13)  # openpyxl 은 16 유효숫자로 쓴다
            self.assertEqual(ws["B22"].font.name, "Calibri"); self.assertEqual(ws["C22"].font.name, "Calisto MT"); self.assertEqual(ws["C22"].number_format, "0.00000_ ")
            self.assertEqual(wb["Rd_dc"]["C40"].value, 10000); self.assertTrue(str(wb["Rd_dc"]["C19"].value).startswith("=PV("))
            self.assertTrue(wb.loaded_theme)
            blocks = dc_blocks(s, C, "RF")
            self.assertEqual([b["key"] for b in blocks], ["block1", "block2", "block3", "block4"]); self.assertEqual(blocks[0]["orient"], "rows")
            # 파이썬 모형(화면 값) vs 엔진 state: 이표 격자 DF 는 같은 정의(폐형식)라 1e-12 안에서 일치해야 한다
            row22 = next(r for r in blocks[2]["rows"] if r["key"] == "row22")["values"]
            self.assertLess(max(abs(a - b) for a, b in zip(row22, s.bootstrap.df["RF"])), 1e-12)
            # 격자 주간 선도(r13) vs 엔진 fwd.disc_per_step(1_discrete_fwd): 분기 안 일정 선도 — 같은 원리.
            # 엔진 격자 시각은 T_ROUND_DIGITS 로 반올림(1/52 → 0.01923077)되어 상대 4e-8 차이(절대 ≤4e-10)가 정상(REVIEWER_SHEET_SPEC §6)
            r13 = next(r for r in blocks[1]["rows"] if r["key"] == "row13")["values"]
            self.assertLess(max(abs(a - b) for a, b in zip(r13, s.fwd.disc_per_step["RF"]["values"])), 1e-9)
        finally:
            shutil.rmtree(base, ignore_errors=True)

    def test_fallback_value_sheets_for_daily(self):
        from openpyxl import load_workbook
        from cb_valuation.step1_curve.io.evidence_writer import export_xlsx, dc_blocks
        base = tempfile.mkdtemp(prefix="cb_rs_")
        try:
            s, C = drive_to_done(base, "REVIEWER_2024", step="daily", horizon_years=2)
            self.assertEqual(s.run.status, "done", s.result.fail_reason)
            ok, why = RS.applicable(s, C); self.assertFalse(ok); self.assertIn("주간·월간만", why)
            path = export_xlsx(s, C, base); ws = load_workbook(path)["Rf_dc"]
            self.assertTrue(str(ws["B2"].value).startswith("값 시트"))
            self.assertNotIn("orient", dc_blocks(s, C, "RF")[0])
        finally:
            shutil.rmtree(base, ignore_errors=True)

    def test_default_profile_uses_value_sheets(self):
        base = tempfile.mkdtemp(prefix="cb_rs_")
        try:
            s, C = drive_to_done(base, "DEFAULT", step="weekly", horizon_years=10)
            self.assertFalse(RS.applicable(s, C)[0])
        finally:
            shutil.rmtree(base, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()


class TestInterpolationChoice(unittest.TestCase):
    """앱 '보간법' 두 선택지(프로필 LINEAR/PCHIP, 사용자 결정 2026-09-08): 검토자 방식 수식 시트 + 국고채 반기(RF_FREQ 2)·회사채 분기(RD_FREQ 4)."""

    def test_linear_and_pchip_profiles(self):
        from openpyxl import load_workbook
        from cb_valuation.step1_curve.io.evidence_writer import export_xlsx, dc_blocks
        for profile in ("LINEAR", "PCHIP"):
            base = tempfile.mkdtemp(prefix="cb_rs_")
            try:
                s, C = drive_to_done(base, profile, step="weekly", horizon_years=10)
                self.assertEqual(s.run.status, "done", s.result.fail_reason)
                self.assertEqual((C.RF_FREQ, C.RD_FREQ), (2, 4))
                ok, why = RS.applicable(s, C); self.assertTrue(ok, why)
                for c in ("RF", "RD"):
                    p = RS.params_from_state(s, C, c); V = RS.evaluate(p); R = p.rows
                    self.assertEqual(p.method, "pchip" if profile == "PCHIP" else "linear")
                    self.assertEqual(p.m, 2 if c == "RF" else 4)
                    self.assertLess(max(abs(a - b) for a, b in zip(V[R["qdf"]], s.bootstrap.df[c])), 1e-12)  # 이표 격자 DF = 엔진
                    self.assertLess(max(abs(a - b) for a, b in zip(V[R["qy"]], s.interp.ytm_on_coupon_grid[c]["values"])), 1e-12)  # 행 18 = 엔진 이표격자 YTM(선형/PCHIP)
                    self.assertLess(max(abs(a - b) for a, b in zip(V[R["ytm"]], s.tree.ytm_on_grid[c]["values"][1:])), 1e-9)  # 행 11(격자 시각 반올림 차이)
                    self.assertIs(V[R["mc1"]][0], True); self.assertIs(V[R["mc2"]][0], True)
                path = export_xlsx(s, C, base); wb = load_workbook(path); rf, rd = wb["Rf_dc"], wb["Rd_dc"]
                self.assertEqual(rf["B16"].value, "HALF-YEAR"); self.assertEqual(rf["C17"].value, "=C16*26"); self.assertEqual(rf["C19"].value, "=C18/2")
                self.assertEqual(rf["B19"].value, "HALF-YEARLY PAYMENT RATE"); self.assertEqual(rf["C13"].value, "=((1+$AB$32)^26)^(1/26)-1")
                self.assertEqual(rd["B16"].value, "QUARTER"); self.assertEqual(rd["C17"].value, "=C16*13"); self.assertEqual(rd["C13"].value, "=((1+$O$30)^13)^(1/13)-1")
                if profile == "PCHIP":
                    self.assertEqual(rf["B8"].value, "PCHIP SLOPE (dY/dSTEP)")
                    self.assertTrue(str(rf["C8"].value).startswith("=IF(SIGN(")); self.assertTrue(str(rf["D8"].value).startswith("=IF(OR(SIGN("))
                    self.assertTrue(str(rf["C11"].value).startswith("=IF(C$10<=$C$4,$C$6,IF(C$10>=$L$4,$L$6,(2*C$41^3"))
                    self.assertEqual(rf["C40"].value, "=IFERROR(MATCH(C$10,$C$4:$L$4,1),1)"); self.assertEqual(rd["B54"].value, "PCHIP SEGMENT")
                    self.assertIn("row8", [r["key"] for r in dc_blocks(s, C, "RF")[0]["rows"]])
                else:
                    self.assertIsNone(rf["B8"].value); self.assertIsNone(rf["C40"].value)
                    self.assertTrue(str(rf["C11"].value).startswith("=IF(C$10<=$C$4,$C$6,IF(C$10>=$L$4,$L$6,INDEX("))
            finally:
                shutil.rmtree(base, ignore_errors=True)

    def test_ui_profiles(self):
        from cb_valuation.step1_curve.app.runner import profiles_info
        vis = [p for p in profiles_info() if p["visible"]]
        self.assertEqual([p["name"] for p in vis], ["LINEAR", "PCHIP"])
        self.assertEqual([p["label"] for p in vis], ["선형 보간", "PCHIP"])
        self.assertTrue(all(p["implemented"] for p in vis))

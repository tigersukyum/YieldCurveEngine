# -*- coding: utf-8 -*-
"""
io/evidence_writer.py — 증빙 번들 작성기 (GRAPH_SPEC §10, docs/XLSX_TEMPLATE.md). state 값 복사만(배율·서식 제외) — 파생값은 노드가 state 에 기록한다.
evidence/<valuation_date>__<curve_set_id>/ 에 EVIDENCE_FILES 01~12 + README_conventions.md + checklist_map.json + evidence.xlsx(필수).
xlsx 의 Rf_dc/Rd_dc 는 XLSX_DC_BLOCKS/XLSX_DC_STYLE(검토자 2024 패키지 배치) 대로 행 지향으로 쓰고, 그 밖의 시트는 XLSX_COLUMNS 열 지향이다.
값은 state 의 double 그대로 쓰고 서식만 입힌다. 위치표(cell_map)는 실제 쓴 범위로 만든다.
"""
from __future__ import annotations
import csv, json, math, os
from ..graph import step1_graph as G
from . import reviewer_sheet as RS

XLSX_NAME = "evidence.xlsx"


def _key(s):
    return f"{s.provenance.valuation_date or 'nodate'}__{s.run.curve_set_id or 'noset'}"


def _csv(path, header, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh); w.writerow(header)
        for r in rows:
            w.writerow(["" if v is None else v for v in r])


def _json(path, obj):
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(json.dumps(json.loads(G.canonical_json(obj)), ensure_ascii=False, indent=1))


def _col(n: int) -> str:
    out = ""
    while n:
        n, r = divmod(n - 1, 26); out = chr(65 + r) + out
    return out


def dc_values(s, C, c: str) -> dict:
    """XLSX_DC_BLOCKS 의 행 라벨(치환 전 원문) → 값 리스트. **state 값 복사만** 한다(PAR_FACE 배율·WEEKS 환산은 표시 배율).
    블록 1·3 = 마디/부트스트랩 격자, 블록 2·4 = 트리 격자 스텝. 정의는 docs/XLSX_TEMPLATE.md §4."""
    knots = s.grid.knots[c]
    ann = dict(zip(s.conv.spot_annual[c]["times"], s.conv.spot_annual[c]["values"]))
    cont = dict(zip(s.conv.spot_cont[c]["times"], s.conv.spot_cont[c]["values"]))
    bt = s.grid.boot_times[c]
    pts = s.bootstrap.points[c]
    par = s.par_check.per_maturity[c]
    times = s.grid.tree.times
    N = len(times) - 1
    fsr = s.fwd_spot_check.rows.get(c, []) if isinstance(s.fwd_spot_check.rows, dict) else []
    W = C.WEEKS_PER_YEAR
    return {
        # 블록 1 (마디)
        "WEEKS": [round(k["t"] * W) for k in knots], "TENOR": [k["tenor"] for k in knots],
        "{rate} - YTM": [k["ytm"] for k in knots], "SPOT RATE": [ann.get(round(k["t"], C.T_ROUND_DIGITS)) for k in knots],
        # 블록 2 (트리 격자)
        "STEP": list(range(1, N + 1)), "t (years)": times[1:], "{rate} - YTM#grid": s.tree.ytm_on_grid[c]["values"][1:],
        "SPOT RATE#grid": s.tree.spot_annual_on_grid[c]["values"][1:], "FORWARD RATE": s.fwd.disc_per_step[c]["values"],
        # 블록 3 (부트스트랩 격자)
        "{period}": [p["n"] for p in pts], "WEEKS#boot": [round(t * W) for t in bt], "YTM - YEARLY": s.interp.ytm_on_coupon_grid[c]["values"],
        "{period} PAYMENT RATE": [p["c"] for p in pts], "PV OF PRINCIPAL": [C.PAR_FACE * p["df"] for p in pts],
        "PV OF BOND": [C.PAR_FACE * r["target"] for r in par], "PVF OF SPOT Rate": [p["df"] for p in pts],
        "SUM OF PVF SPOT JUST PRIOR TO": [p["sum_df_prior"] for p in pts], "{period} SPOT Rate": [p["spot_pp"] for p in pts],
        "SPOT Rate -Yearly": [ann.get(t) for t in bt], "SPOT Rate -Continuous": [cont.get(t) for t in bt],
        "FORWARD Rate - {period}": s.fwd.boot_fwd_pp[c]["values"], "FORWARD Rate - Yearly": s.fwd.boot_fwd_annual[c]["values"],
        "MODEL CHECK (PAR REPRICE)": [C.PAR_FACE * r["price"] for r in par], "PAR RESIDUAL": [r["residual"] for r in par],
        # 블록 4 (트리 격자; 검토자 r32~r38 의미)
        "STEP SPOT RATE": s.fwd.spot_per_step[c]["values"], "FORMULA I": s.fwd.growth_step[c], "FORMULA II -CUMM": s.fwd.growth_cum_prev[c],
        "STEP FORWARD RATE": s.fwd.disc_per_step[c]["values"], "PVF OF FORWARD RATE": s.fwd.df_step[c],
        "MODEL CHECK (PROD DF_FWD - DF_SPOT)": [r["diff_prod"] for r in fsr] if fsr else [None] * N,
        "MODEL CHECK (PV)": [C.PAR_FACE * d for d in s.fwd.df_backward[c]],
    }


def dc_blocks(s, C, c: str) -> list:
    """블록 구조 [{title, rows:[{label, key, fmt, values}]}] — xlsx 작성기와 화면(/api/blocks)이 같은 함수를 쓴다.
    검토자 수식 시트가 적용되는 실행(reviewer_sheet.applicable: 검토자 방식 상수 + 주간/월간 격자, XLSX_FORMULA_SHEETS)은 그 시트의 행(파이썬 모형 값)을,
    아니면 XLSX_DC_BLOCKS(값 시트)를 돌려준다. 어느 쪽인지는 각 블록의 `orient`('rows' = 검토자 배치)로 화면이 안다."""
    if C.XLSX_FORMULA_SHEETS and RS.applicable(s, C)[0]:
        return RS.blocks(RS.params_from_state(s, C, c))
    m = C.RF_FREQ if c == "RF" else C.RD_FREQ
    rate = "RISK FREE RATE" if c == "RF" else "RISKY RATE"
    period = {1: "YEAR", 2: "HALF-YEAR", 4: "QUARTER", 12: "MONTH"}.get(m, f"1/{m}Y")
    vals = dc_values(s, C, c)
    out = []
    for bi, (title, rows) in enumerate(C.XLSX_DC_BLOCKS, start=1):
        block = {"index": bi, "title": title.format(rate=rate, period=period), "rows": []}
        for label, _src, fmt in rows:
            key = label
            if bi == 2 and label in ("{rate} - YTM", "SPOT RATE"): key = label + "#grid"
            if bi == 3 and label == "WEEKS": key = "WEEKS#boot"
            block["rows"].append({"label": label.format(rate=rate, period=period), "key": key, "fmt": fmt, "values": list(vals.get(key, []))})
        out.append(block)
    return out


def _write_dc_sheet(ws, s, C, c: str) -> dict:
    from openpyxl.styles import Font
    st = C.XLSX_DC_STYLE
    ws[st["title_cell"]] = "무위험이자율" if c == "RF" else "위험이자율"; ws[st["title_cell"]].font = Font(bold=True)
    ws[st["date_cell"]] = s.provenance.valuation_date
    ws.column_dimensions[st["label_col"]].width = st["width_label"]
    ws.freeze_panes = st["freeze_panes"]
    r = 3
    ranges = {}
    first_col = ord(st["first_data_col"]) - 64
    max_col = first_col
    for block in dc_blocks(s, C, c):
        bi = block["index"]
        ws.cell(r, 2, block["title"]).font = Font(bold=st["block_title_bold"])
        r0 = r; r += 1
        if bi == 1:
            ws.cell(r, 2, s.provenance.curve_date)  # REV Rf_dc!B5 = 고시일
        for row in block["rows"]:
            data = row["values"]
            ws.cell(r, 2, row["label"])
            for j, v in enumerate(data):
                cell = ws.cell(r, first_col + j, v)
                cell.number_format = row["fmt"]
            max_col = max(max_col, first_col + max(len(data), 1) - 1)
            ranges[f"{c}:{row['key']}"] = f"{st['first_data_col']}{r}:{_col(first_col + max(len(data), 1) - 1)}{r}"
            r += 1
        ranges[f"{c}:block{bi}"] = f"B{r0}:{_col(max_col)}{r - 1}"
        r += 1
    for j in range(first_col, max_col + 1):
        ws.column_dimensions[_col(j)].width = st["width_data"]
    return ranges


PAR_CHECK_LAYOUT = {  # 검토자 FY25 패키지 '검증' 시트 B13:I29 'Par 검증' 블록의 배치·서식(값 없음): 열 B 라벨, C Date, D/E 일수계산, F/G Rate, H/I Par 검증
    "width_B": 28.09765625, "width_C": 13.8984375, "font": "맑은 고딕", "size": 11, "header_tint": -0.0499893185216834,
    "nf_date": "mm-dd-yy", "nf_act": "0.0000", "nf_pct": "0.00%;[Red]\\-0.00%;\\-", "nf_check": "0.00",
    "tables": (("RF", "국고채", "Rf_dc"), ("RD", "회사채", "Rd_dc")),
}


def _write_par_check_formula(ws, s, C, params: dict) -> str:
    """PAR_CHECK 시트를 검토자 '검증' 시트의 Par 검증 블록처럼: 국고채(이자 6개월)·회사채(이자 3개월) 표 2개, 행 수 = CURVE_HORIZON_Y × m(산출 기간 최대),
    Date = 평가일부터 EOMONTH 체인, 30/360·Act/365F 일수계산, YTM·현물은 Rf_dc/Rd_dc 이표 격자(행 16 기간번호 → 행 18 YTM, DF 행)에서 HLOOKUP,
    Par 검증 H = YTM 항등식(Σ c/m /(1+y/m)^(t·m) + 1/(1+y/m)^(t·m) = 1), I = 현물 DF 재가격(Σ c/m·(1+g_k)^−t_k + (1+g_n)^−t_n = 1; g = (1/DF)^(1/t)−1).
    시트 범위 밖(산출 기간 초과) 행은 IFERROR → 0 이라 검증값 1 로 표시된다(원본 패키지와 같은 동작). 반환: 국고채 표 범위(cell_map C36)."""
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Color
    from datetime import datetime
    L = PAR_CHECK_LAYOUT
    thin = Side(style="thin")
    B_all = Border(left=thin, right=thin, top=thin, bottom=thin); B_lr = Border(left=thin, right=thin)
    B_lrb = Border(left=thin, right=thin, bottom=thin); B_rtb = Border(right=thin, top=thin, bottom=thin)
    F = Font(name=L["font"], size=L["size"]); FB = Font(name=L["font"], size=L["size"], bold=True)
    gray = PatternFill(patternType="solid", fgColor=Color(rgb=RS.DESIGN["steel"]) if C.XLSX_PALETTE == "design" else Color(theme=0, tint=L["header_tint"]))  # 헤더·평가일 행: 디자인 팔레트 steel(원본은 연회색)
    center = Alignment(horizontal="center")
    ws.column_dimensions["B"].width = L["width_B"]; ws.column_dimensions["C"].width = L["width_C"]
    vd = datetime.fromisoformat(s.provenance.valuation_date)

    def put(r, c, v=None, font=F, fill=None, border=None, nf=None, align=None):
        cell = ws.cell(r, c)
        if v is not None:
            cell.value = v
        cell.font = font
        if fill is not None: cell.fill = fill
        if border is not None: cell.border = border
        if nf is not None: cell.number_format = nf
        if align is not None: cell.alignment = align
        return cell
    r = 2; first_range = None
    for cid, name, sheet in L["tables"]:
        p = params[cid]; m = p.m; months = 12 // m; n_rows = int(round(C.CURVE_HORIZON_Y * m)); Lq = p.Lq
        df_row = p.rows["qdf"]; ytm_row = p.rows["qy"]; idx_row = p.rows["q"]
        put(r, 2, f"Par 검증 - {name} (이자지급 {months}개월, {sheet})", font=FB)
        hdr, base, first = r + 1, r + 2, r + 3
        for c_i, text in ((3, "Date"), (4, "Daycount convention"), (6, "Rate"), (8, "Par 검증")):
            put(hdr, c_i, text, fill=gray, border=B_all, align=center)
        for c_i in (5, 7, 9):
            put(hdr, c_i, border=B_rtb)
        put(base, 2, "평가일", font=FB, fill=gray, border=B_all, align=center)
        put(base, 3, vd, font=FB, fill=gray, border=B_rtb, nf=L["nf_date"], align=center)
        for c_i, text in ((4, "30/360"), (5, "Act/365F"), (6, "YTM"), (7, "Spot"), (8, "1확인")):
            put(base, c_i, text, font=FB, fill=gray, border=B_all, align=center)
        put(base, 9, border=B_rtb)
        for k in range(1, n_rows + 1):
            rr = first + k - 1; prev = base if k == 1 else rr - 1
            put(rr, 2, "쿠폰지급일" if k == 1 else None, border=B_all if k == 1 else (B_lrb if k == n_rows else B_lr), align=center if k == 1 else None)
            put(rr, 3, f"=EOMONTH(C{prev},{months})", border=B_all, nf=L["nf_date"])
            put(rr, 4, f"=YEARFRAC($C${base},C{rr},0)", border=B_all)
            put(rr, 5, f"=(C{rr}-$C${base})/365", border=B_all, nf=L["nf_act"])
            put(rr, 6, f"=IFERROR(HLOOKUP($D{rr}*{m},{sheet}!$C${idx_row}:${Lq}${ytm_row},{ytm_row - idx_row + 1},FALSE),0)", border=B_all, nf=L["nf_pct"])
            put(rr, 7, f"=IFERROR((1/HLOOKUP($D{rr}*{m},{sheet}!$C${idx_row}:${Lq}${df_row},{df_row - idx_row + 1},FALSE))^(1/$D{rr})-1,0)", border=B_all, nf=L["nf_pct"])
            put(rr, 8, f"=SUMPRODUCT(($F{rr}/{m})/(1+$F{rr}/{m})^($D${first}:D{rr}*{m}))+1/(1+$F{rr}/{m})^(D{rr}*{m})", border=B_all, nf=L["nf_check"])
            put(rr, 9, f"=SUMPRODUCT(($F{rr}/{m})/(1+$G${first}:G{rr})^($D${first}:D{rr}))+1/(1+G{rr})^D{rr}", border=B_all, nf=L["nf_check"])
        # 원본과 같은 병합: 헤더 Daycount(D:E)·Rate(F:G)·Par 검증(H:I), 평가일 행 '1확인'(H:I), 쿠폰지급일 라벨(B 첫 행~마지막 행)
        for c1, c2 in ((4, 5), (6, 7), (8, 9)):
            ws.merge_cells(start_row=hdr, start_column=c1, end_row=hdr, end_column=c2)
        ws.merge_cells(start_row=base, start_column=8, end_row=base, end_column=9)
        ws.merge_cells(start_row=first, start_column=2, end_row=first + n_rows - 1, end_column=2)
        rng = f"B{r}:I{first + n_rows - 1}"
        first_range = first_range or rng
        r = first + n_rows + 2
    return first_range


def _write_table(ws, header, rows):
    from openpyxl.styles import Font
    for j, h in enumerate(header, start=1):
        ws.cell(1, j, h).font = Font(bold=True)
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v if not (isinstance(v, float) and not math.isfinite(v)) else str(v))
    return f"A1:{_col(max(len(header), 1))}{len(rows) + 1}"


def _rows_used(s, C):
    out = []
    for c in C.CURVE_IDS:
        r = s.rows[c.lower()]
        if r:
            out.append([c, r["row_index"], r["label_raw"], r.get("block"), r.get("block_requested"), r.get("rating"), r.get("notch"), f"nominal_m{C.RF_FREQ if c == 'RF' else C.RD_FREQ}"] + [r["ytm_pct"][t] for t in C.TENOR_LABELS])
    return out


def _flag_rows(s):
    return [[f["severity"], f["code"], f.get("curve"), json.dumps(f.get("value"), ensure_ascii=False, default=str), f.get("threshold"), f.get("detail")]
            for lst in (s.sanity.fail, s.sanity.approval_required, s.sanity.warn) for f in lst]


def write_xlsx(path: str, s, C):
    """XLSX_SHEETS 순서로 통합문서 작성(Rf_dc/Rd_dc 는 XLSX_DC_BLOCKS). 반환 (dc 범위, 시트 범위). openpyxl 없으면 ImportError."""
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active)
    ws = {name: wb.create_sheet(name) for name in C.XLSX_SHEETS}
    rng, ranges = {}, {}
    rows_used = _rows_used(s, C)
    rng["INPUT_RAW"] = _write_table(ws["INPUT_RAW"], ["row_index", "kind", "group", "label"] + C.TENOR_LABELS, [[r["row_index"], r["kind_raw"], r["group_raw"], r["label_raw"]] + [r["ytm_pct"][t] for t in C.TENOR_LABELS] for r in s.input.rows])
    rng["ROWS_USED"] = _write_table(ws["ROWS_USED"], ["curve", "row_index", "label", "block", "block_requested", "rating", "notch", "basis"] + C.TENOR_LABELS, rows_used)
    prov = [(k, json.dumps(v, ensure_ascii=False, default=str) if isinstance(v, (dict, list)) else v) for k, v in json.loads(G.canonical_json(s.provenance)).items()]
    rng["PROVENANCE"] = _write_table(ws["PROVENANCE"], ["field", "value"], prov)
    rng["CONVENTIONS"] = _write_table(ws["CONVENTIONS"], ["constant", "value"], [(k, json.dumps(v, ensure_ascii=False, default=str)) for k, v in sorted(C.items().items())])
    formula_ok, formula_why = RS.applicable(s, C)
    use_formula = formula_ok and C.XLSX_FORMULA_SHEETS
    params = {c: RS.params_from_state(s, C, c) for c in C.CURVE_IDS} if use_formula else {}
    for c, name in (("RF", "Rf_dc"), ("RD", "Rd_dc")):
        if use_formula:  # 검토자 시트를 살아있는 수식·원본 배치로(docs/REVIEWER_SHEET_SPEC.md); 색은 XLSX_PALETTE
            ranges.update(RS.write_sheet(ws[name], wb, params[c], palette=C.XLSX_PALETTE))
        else:  # 값 시트(XLSX_DC_BLOCKS); 사유를 B2 에 남긴다
            ranges.update(_write_dc_sheet(ws[name], s, C, c))
            ws[name]["B2"] = f"값 시트(수식 시트 미적용: {formula_why if not formula_ok else 'XLSX_FORMULA_SHEETS=False'})"
    if use_formula:  # 검토자 '검증' 시트의 Par 검증 블록 배치·서식으로, Rf_dc/Rd_dc 에서 수식으로 가져와 검증(사용자 요청 2026-09-08)
        rng["PAR_CHECK"] = _write_par_check_formula(ws["PAR_CHECK"], s, C, params)
    else:
        rng["PAR_CHECK"] = _write_table(ws["PAR_CHECK"], C.XLSX_COLUMNS["PAR_CHECK"], [[r["curve"], r["t"], r["n"], r["price"], r["target"], r["residual"], r["residual_x_face"]] for c in C.CURVE_IDS for r in s.par_check.per_maturity[c]])
    fs_rows = [[c, r["step"], r["t"], r["prod_df_fwd"], r["df_spot"], r["diff_log"], r["diff_prod"]] for c in C.CURVE_IDS for r in (s.fwd_spot_check.rows.get(c, []) if isinstance(s.fwd_spot_check.rows, dict) else [])]
    rng["FWD_SPOT_CHECK"] = _write_table(ws["FWD_SPOT_CHECK"], C.XLSX_COLUMNS["FWD_SPOT_CHECK"], fs_rows)
    rng["SENSITIVITY"] = _write_table(ws["SENSITIVITY"], ["method", "space", "curve", "max_rel_df_diff", "note"], [[r["method"], r["space"], r["curve"], r["max_rel_df_diff"], r.get("note")] for r in s.sensitivity.table] + [["freq_alt", "", c, json.dumps(v, ensure_ascii=False), ""] for c, v in dict(s.sensitivity.freq_alt).items()])
    hl = json.loads(G.canonical_json(s.headline))
    rng["HEADLINE"] = _write_table(ws["HEADLINE"], ["field", "value"], [(k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v) for k, v in hl.items()])
    rng["FLAGS"] = _write_table(ws["FLAGS"], ["severity", "code", "curve", "value", "threshold", "detail"], _flag_rows(s))
    ap_rows = [[k, a.requested_at, a.snapshot_sha256, json.dumps([f["code"] for f in a.flags_seen], ensure_ascii=False), a.decision, a.approver, a.timestamp, a.comment, json.dumps(a.acknowledged_codes, ensure_ascii=False)]
               for k, a in (("input", s.approval_input), ("exception", s.approval_exception), ("curve", s.approval_curve))]
    rng["APPROVALS"] = _write_table(ws["APPROVALS"], C.XLSX_COLUMNS["APPROVALS"], ap_rows)
    rng["RUN_PATH"] = _write_table(ws["RUN_PATH"], C.XLSX_COLUMNS["RUN_PATH"], [[i + 1] + list(p) for i, p in enumerate(s.run.path)])
    wb.save(path)
    return ranges, rng


def evidence_zip_bytes(s, base_dir: str) -> bytes:
    """증빙 번들 폴더(state.export.dir) 전체를 zip 바이트로(화면 '증빙 zip 내려받기'; 브라우저 실행에서는 파일이 메모리에만 있어 이 경로가 유일한 반출)."""
    import io as _io, zipfile
    d = os.path.join(base_dir, s.export.dir)
    buf = _io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for root, _dirs, files in os.walk(d):
            for f in sorted(files):
                p = os.path.join(root, f)
                z.write(p, os.path.relpath(p, os.path.dirname(d)).replace("\\", "/"))
    return buf.getvalue()


def export_xlsx(s, C, base_dir: str) -> str:
    """화면 '엑셀 내려받기': 현재 state 로 같은 형식의 통합문서를 exports/ 에 쓴다(증빙 번들과 별개, 승인 불필요). 반환 경로."""
    d = os.path.join(base_dir, "exports"); os.makedirs(d, exist_ok=True)
    from datetime import datetime, timezone
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    path = os.path.join(d, f"curve_{_key(s)}_{stamp}.xlsx")
    write_xlsx(path, s, C)
    return path


def write_bundle(s, C, base_dir: str) -> dict:
    key = _key(s)
    d = os.path.join(base_dir, "evidence", key); os.makedirs(d, exist_ok=True)
    files, warnings, cell_map = [], [], {}
    def add(name): files.append({"path": name, "abs": os.path.join(d, name)})
    ef = {k: f"{k}_{n}.{e}" for k, (n, e) in C.EVIDENCE_FILES.items()}
    # 01 raw matrix
    _csv(os.path.join(d, ef["01"]), ["row_index", "kind", "group", "label"] + C.TENOR_LABELS,
         [[r["row_index"], r["kind_raw"], r["group_raw"], r["label_raw"]] + [r["ytm_pct"][t] for t in C.TENOR_LABELS] for r in s.input.rows]); add(ef["01"])
    # 02 rows used
    rows_used = _rows_used(s, C)
    _csv(os.path.join(d, ef["02"]), ["curve", "row_index", "label", "block", "block_requested", "rating", "notch", "basis"] + C.TENOR_LABELS, rows_used); add(ef["02"])
    # 03 knots
    _csv(os.path.join(d, ef["03"]), ["curve", "tenor", "t", "ytm[nominal]"], [[c, k["tenor"], k["t"], k["ytm"]] for c in C.CURVE_IDS for k in s.grid.knots[c]]); add(ef["03"])
    # 04 bootstrap
    boot_rows = []
    for c in C.CURVE_IDS:
        m = C.RF_FREQ if c == "RF" else C.RD_FREQ
        for p in s.bootstrap.points[c]:
            boot_rows.append([c, p["n"], p["t"], p["c"], p["price_target"], p["spot_pp"], p["df"], p["denominator"], f"per_period_m{m}"])
    _csv(os.path.join(d, ef["04"]), ["curve", "n", "t", "c[per_period]", "price_target", "spot_pp[per_period]", "DF", "denominator", "basis"], boot_rows); add(ef["04"])
    # 05 spot table
    spot_rows = []
    for c in C.CURVE_IDS:
        for t, spp, z, r, df in zip(s.grid.boot_times[c], s.bootstrap.spot_pp[c]["values"], s.conv.spot_annual[c]["values"], s.conv.spot_cont[c]["values"], s.bootstrap.df[c]):
            spot_rows.append([c, t, spp, z, r, df])
    _csv(os.path.join(d, ef["05"]), ["curve", "t", "spot_pp[per_period]", "spot_annual[annual_eff]", "spot_cont[continuous]", "DF"], spot_rows); add(ef["05"])
    # 06 tree grid
    tg = []
    times = s.grid.tree.times
    for c in C.CURVE_IDS:
        za, zc, dfs = s.tree.spot_annual_on_grid[c]["values"], s.tree.spot_cont_on_grid[c]["values"], s.tree.df_spot_on_grid[c]
        f, F, Fa, ds, da, dc = (s.fwd.cont_on_grid[c]["values"], s.fwd.disc_per_step[c]["values"], s.fwd.disc_annual_eff[c]["values"], s.fwd.df_step[c], s.fwd.df_step_alt[c], s.fwd.df_cum[c])
        flags = {i: "left_flat" for i in s.tree.extrap_left_flat_steps[c]}; flags.update({i: "right" for i in s.tree.extrap_right_steps[c]})
        for i in range(1, len(times)):
            tg.append([c, i, times[i], s.grid.tree.dt[i - 1], za[i], zc[i], dfs[i], f[i - 1], F[i - 1], Fa[i - 1], ds[i - 1], da[i - 1], dc[i - 1], flags.get(i, "")])
    _csv(os.path.join(d, ef["06"]), ["curve", "step", "t", "dt", "D-SPOT[annual_eff]", "C-SPOT[continuous]", "DF", "C-FWD[continuous]", "F_step[per_step_simple]", "F_annual[annual_eff]", "DF_step", "DF_step_alt", "DF_cum", "extrap_flag"], tg); add(ef["06"])
    # 07 par residuals
    _csv(os.path.join(d, ef["07"]), C.XLSX_COLUMNS["PAR_CHECK"], [[r["curve"], r["t"], r["n"], r["price"], r["target"], r["residual"], r["residual_x_face"]] for c in C.CURVE_IDS for r in s.par_check.per_maturity[c]]); add(ef["07"])
    # 08 fwd-spot sample (md)
    smp = s.fwd_spot_check.sample or {}
    with open(os.path.join(d, ef["08"]), "w", encoding="utf-8", newline="\n") as fh:
        fh.write("# 감사인 Q11 — Π DF_fwd = DF_spot 예시 1건\n\n")
        fh.write(f"- 커브 {smp.get('curve')} / step {smp.get('step')} / t = {smp.get('t')}\n- Π df_step = {smp.get('prod_df_fwd')!r}\n- DF_spot(t) = {smp.get('df_spot')!r}\n")
        fh.write(f"- 차이(곱) = {smp.get('diff_prod')!r}, 차이(로그) = {smp.get('diff_log')!r}\n- 전 격자점 최대: 로그 {s.fwd_spot_check.max_abs_err_log!r}, 곱 {s.fwd_spot_check.max_abs_err_prod!r} (게이트 TOL_FWD_SPOT_FAIL)\n")
    add(ef["08"])
    # 09 flags
    flag_rows = _flag_rows(s)
    _csv(os.path.join(d, ef["09"]), ["severity", "code", "curve", "value", "threshold", "detail"], flag_rows); add(ef["09"])
    # 10 headline
    _json(os.path.join(d, ef["10"]), dict(s.headline)); add(ef["10"])
    # 11 sensitivity
    _csv(os.path.join(d, ef["11"]), ["method", "space", "curve", "max_rel_df_diff", "note"], [[r["method"], r["space"], r["curve"], r["max_rel_df_diff"], r.get("note")] for r in s.sensitivity.table]); add(ef["11"])
    # 12 approvals
    _json(os.path.join(d, ef["12"]), {k: dict(s[f"approval_{k}"]) for k in ("input", "exception", "curve")}); add(ef["12"])
    # README_conventions.md
    stmt = "\n".join(f"- **{k}** = `{getattr(C, k)}` — {v}" for k, v in C.CONVENTION_REASONS.items() if hasattr(C, k)) + \
           f"\n- **EXTRAP** = EXTRAP_LEFT `{C.EXTRAP_LEFT}` / EXTRAP_RIGHT `{C.EXTRAP_RIGHT}` — {C.CONVENTION_REASONS['EXTRAP']}"
    mc = s.provenance.method_choice
    with open(os.path.join(d, "README_conventions.md"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write(f"# 관례 선언문 — {key} (프로필 {s.run.profile})\n\n## 프로필(사용자 선택, PROFILE_SELECTION={C.PROFILE_SELECTION})\n- {mc.profile}: {C.PROFILE_DESCRIPTIONS.get(mc.profile, '')}\n- 선택자 {mc.chosen_by}, 시각 {mc.chosen_at}; 보간 {mc.interp_method} / 공간 {mc.interp_space_grid}\n\n")
        fh.write("## 관례(CONVENTION_REASONS)\n" + stmt + "\n\n")
        fh.write(f"## INTERP_METHOD\n- 방법 `{C.INTERP_METHOD}` / 트리 격자 공간 `{C.INTERP_SPACE_GRID}` / 이표격자 공간 `{C.INTERP_SPACE_PRE}`(선형) — 한공회 §3.7.3.6 공시 항목(KICPA_INTERP_DISCLOSURE)\n\n")
        fh.write(f"## NODE_DISCOUNT_CONV\n- `{C.NODE_DISCOUNT_CONV}` — {C.CONVENTION_REASONS['NODE_DISCOUNT_CONV']}\n\n## COUPON_CONV\n- `{C.COUPON_CONV}` — {C.CONVENTION_REASONS['COUPON_CONV']}\n\n")
        fh.write(f"## 상수 지문\n- constants_fingerprint `{s.run.constants_fingerprint}`\n\n## 상수 전부\n")
        for k, v in sorted(C.items().items()):
            fh.write(f"- {k} = `{json.dumps(v, ensure_ascii=False, default=str)[:300]}`\n")
    add("README_conventions.md")
    # xlsx (필수) — 같은 작성기를 화면 '엑셀 내려받기'(export_xlsx) 도 쓴다
    xlsx_written, xlsx_path, ranges, rng = False, None, {}, {}
    try:
        xlsx_path = os.path.join(d, XLSX_NAME)
        ranges, rng = write_xlsx(xlsx_path, s, C)
        xlsx_written = True; add(XLSX_NAME)
    except ImportError:
        warnings.append("XLSX_SKIPPED: openpyxl 없음"); xlsx_path = None
    X = XLSX_NAME
    n06 = len(tg) + 1
    cell_map = {
        "Q1": f"{X}!Rf_dc!{ranges.get('RF:block3', 'B1')}", "Q8": "README_conventions.md!-!NODE_DISCOUNT_CONV",
        "Q9_FWD_INPUT": f"{ef['06']}!-!A1:N{n06}", "Q10_1": f"{ef['05']}!-!A1:F{len(spot_rows) + 1}", "Q10_2": "README_conventions.md!-!COUPON_CONV",
        "Q11": f"{ef['08']}!-!sample", "Q12": f"{X}!PROVENANCE!{rng.get('PROVENANCE', 'A1')}", "Q13": f"{ef['02']}!-!A1:X{len(rows_used) + 1}",
        "C33": f"{X}!Rf_dc!{ranges.get('RF:block4', 'B1')}", "C34": f"{ef['06']}!-!DF_step", "C35": f"{X}!Rf_dc!{ranges.get('RF:block3', 'B1')}",
        "C36": f"{X}!PAR_CHECK!{rng.get('PAR_CHECK', 'A1')}", "C37": "README_conventions.md!-!INTERP_METHOD", "C38": f"{X}!Rd_dc!{ranges.get('RD:block3', 'B1')}",
        "C39": f"{X}!Rd_dc!{ranges.get('RD:block1', 'B1')}", "C40": f"{X}!Rd_dc!{ranges.get('RD:block2', 'B1')}", "C41": f"{ef['06']}!-!A1:N{n06}",
        "C43": f"{ef['05']}!-!A1:F{len(spot_rows) + 1}", "C44": f"{ef['10']}!-!rf_ytm_remaining", "C48": f"{ef['05']}!-!spot_annual[annual_eff]",
        "HEADLINE_RF": f"{ef['10']}!-!rf_ytm_remaining", "HEADLINE_RD": f"{ef['10']}!-!rd_ytm_remaining", "HEADLINE_RATING": f"{ef['10']}!-!rating_applied",
        "HEADLINE_BLOCK": f"{ef['10']}!-!block_applied", "KICPA_INTERP_DISCLOSURE": "README_conventions.md!-!INTERP_METHOD",
        "RUN_PATH": f"{X}!RUN_PATH!{rng.get('RUN_PATH', 'A1')}", "APPROVALS": f"{ef['12']}!-!$",
    }
    _json(os.path.join(d, "checklist_map.json"), cell_map); add("checklist_map.json")
    return {"dir": os.path.relpath(d, base_dir).replace("\\", "/"), "files": files, "cell_map": cell_map, "conventions_statement": stmt,
            "xlsx_written": xlsx_written, "xlsx_path": (os.path.relpath(xlsx_path, base_dir).replace("\\", "/") if xlsx_path else None), "warnings": warnings}

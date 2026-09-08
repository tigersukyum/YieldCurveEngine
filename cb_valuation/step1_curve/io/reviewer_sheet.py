# -*- coding: utf-8 -*-
"""
io/reviewer_sheet.py — 내부 검토자(REV 8521) `Rf_dc`/`Rd_dc` 시트를 **살아있는 엑셀 수식**과 **원본 서식**으로 재현하는 작성기
+ 같은 수식을 파이썬으로 평가하는 시트 모형(화면 표시·검증용). 명세: docs/REVIEWER_SHEET_SPEC.md (행별 수식·서식·역산 근거).

원칙
- 통합문서의 숫자 셀은 입력(테너 주수·YTM·액면 10000·인덱스 시작값)만 값이고 나머지는 전부 수식이다. 파이썬 모형은 같은 수식을 같은 순서로 평가한다
  (정수 거듭제곱은 Excel 과 같은 right-to-left 이진 거듭제곱 `_xpow`).
- 서식은 reviewer_dc_layout.json(원본에서 추출한 서식 런 + 테마 XML; 값 없음)에서 가져온다. 열 수가 원본(520주·40분기·12테너)과 달라도
  열의 성격(첫 열·분기말·분기 첫 주·테너 주·마지막 열)으로 원본 열을 골라 서식을 입힌다(_orig_col).
- 적용 조건(applicable): 검토자 방식 상수(TREE_FWD_RULE=piecewise_quarter_step, NODE_DISCOUNT_CONV=1_discrete_fwd, 선형 YTM 보간, log_df 격자)
  이고 격자 스텝(주간/월간)이 이표주기·테너와 정수로 맞을 때. 아니면 evidence_writer 의 값 시트(_write_dc_sheet)로 내려간다.
- 작성기는 계산하지 않는다(값 셀은 입력뿐). 모형(evaluate)은 화면·테스트 전용이며 state 에 쓰지 않는다.
"""
from __future__ import annotations
import json, math, os
from datetime import datetime

_HERE = os.path.dirname(os.path.abspath(__file__))
_LAYOUT = None
FACE = 10000  # 검토자 시트 액면(원본 C21/C40 리터럴) — 표시 배율일 뿐 커브 값과 무관
SHEET_NAMES = {"RF": "Rf_dc", "RD": "Rd_dc"}
STEP_WORDS = {52: "WEEK", 12: "MONTH"}
PERIOD_WORDS = {4: "QUARTER", 2: "HALF-YEAR", 12: "MONTH", 1: "YEAR"}


def layout() -> dict:
    global _LAYOUT
    if _LAYOUT is None:
        with open(os.path.join(_HERE, "reviewer_dc_layout.json"), encoding="utf-8") as fh:
            _LAYOUT = json.load(fh)
    return _LAYOUT


def theme_bytes() -> bytes:
    with open(os.path.join(_HERE, "reviewer_theme1.xml"), "rb") as fh:
        return fh.read()


def col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26); s = chr(65 + r) + s
    return s


def _xpow(x: float, n: int) -> float:
    """Excel `x^n`(정수 n) 의미론: right-to-left 이진 거듭제곱(검토자 시트 (1+q)^4=(x²)², (1+r)^52 비트 일치)."""
    if n < 0:
        return 1.0 / _xpow(x, -n)
    r, b = 1.0, x
    while n:
        if n & 1:
            r *= b
        b *= b; n >>= 1
    return r


def _pow(x: float, e: float) -> float:
    """Excel `x^e`: e 가 정수면 _xpow, 아니면 C pow."""
    if float(e).is_integer():
        return _xpow(x, int(e))
    return x ** e


# ----------------------------------------------------------------------------- 행 번호표(원본 배치 그대로)
ROWS = {  # slope=PCHIP 마디 기울기(행 8, PCHIP 일 때만), h_seg/h_s = PCHIP 보조행(원본 마지막 행 뒤; PCHIP 일 때만)
    "RF": dict(idx=10, ytm=11, spot=12, fwd=13, qtitle=15, q=16, qw=17, qy=18, qc=19, qpp=20, qpb=21, qdf=22, qs=23, qsp=24, qsy=25, qfq=26, qfy=27, qmc=28,
               mk=30, widx=31, ws=32, f1=33, f2=34, wf=35, pvf=36, mc1=37, mc2=38, last=39, slope=8, h_seg=40, h_s=41),
    "RD": dict(idx=10, ytm=11, spot=12, fwd=13, qtitle=15, q=16, qw=17, qy=18, qpb=19, qdf=20, qs=21, qsp=22, qsy=23, qfq=24, qfy=25, qmc=26,
               mk=28, widx=29, ws=30, f1=31, f2=32, wf=33, wfy=34, mc1=35, mc2=36, pvf=37,
               wbtitle=39, face=40, widx2=41, wc=42, wdf=43, wsum=44, wsp=45, wmc=46, wf1=47, wf2=48, wfw=49, wpvf=50, wmc1=51, wmc2=52, last=53, slope=8, h_seg=54, h_s=55),
}
METHODS = ("linear", "pchip")  # 마디 YTM 보간(행 11): 선형 = 원본 수식, PCHIP = 행 8 기울기(reference/interp_ref.pchip_slopes 와 같은 식) + 보조행 + 3차 Hermite
# 행 성격(서식 열 대응 규칙): fixed=열 그대로(≤R), tenor=테너 열, weekly=격자 스텝 열, quarterly=이표 격자 열
_KINDS = {
    "RF": {**{r: "fixed" for r in (1, 2, 3, 8, 14, 39)}, **{r: "tenor" for r in (4, 5, 6, 7, 15, 29)}, **{r: "weekly" for r in (9, 10, 11, 12, 13, *range(30, 39))},
           **{r: "quarterly" for r in range(16, 29)}},
    "RD": {**{r: "fixed" for r in (1, 2, 3, 8, 14)}, **{r: "tenor" for r in (4, 5, 6, 7, 15, 27)}, **{r: "weekly" for r in (9, 10, 11, 12, 13, *range(28, 54))},
           **{r: "quarterly" for r in range(16, 27)}},
}
# 블록(화면·cell_map): (키, 제목 행, 데이터 행들, 열 종류)
_BLOCKS = {
    "RF": [("block1", 3, [4, 5, 6, 7], "tenor"), ("block2", 9, [10, 11, 12, 13], "weekly"), ("block3", 15, list(range(16, 29)), "quarterly"), ("block4", 30, list(range(31, 39)), "weekly")],
    "RD": [("block1", 3, [4, 5, 6, 7], "tenor"), ("block2", 9, [10, 11, 12, 13], "weekly"), ("block3", 15, list(range(16, 27)), "quarterly"), ("block4", 28, list(range(29, 38)), "weekly"),
           ("block5", 39, list(range(41, 53)), "weekly")],
}


class Params:
    """시트 하나의 입력: kind(RF|RD), P(연간 스텝 수), m(이표 주기/년), N(스텝 수), tenors[{label, years, step, ytm}], curve_date, header."""

    def __init__(self, kind, P, m, N, tenors, curve_date=None, header="", face=FACE, method="linear"):
        self.kind, self.P, self.m, self.N, self.tenors, self.curve_date, self.header, self.face = kind, int(P), int(m), int(N), tenors, curve_date, header, float(face)
        self.method = method
        if method not in METHODS:
            raise ValueError(f"보간법 {method!r} 은 수식 시트에 없음(METHODS={METHODS})")
        if method == "pchip" and len(tenors) < 3:
            raise ValueError("PCHIP 수식 시트는 마디 3개 이상 필요")
        if self.P % self.m:
            raise ValueError(f"P={P} 가 m={m} 의 배수가 아님")
        self.spq = self.P // self.m
        if self.N % self.spq:
            raise ValueError(f"N={N} 이 스텝/이표기간={self.spq} 의 배수가 아님")
        self.Nq = self.N // self.spq
        self.nT = len(tenors)
        self.tenor_steps = {int(t["step"]) for t in tenors}
        self.rows = ROWS[kind]
        self.sheet = SHEET_NAMES[kind]

    @property
    def Lw(self): return col(self.N + 2)      # 마지막 격자 열
    @property
    def Lq(self): return col(self.Nq + 2)     # 마지막 이표격자 열
    @property
    def Lt(self): return col(self.nT + 2)     # 마지막 테너 열
    @property
    def Xc(self): return col(self.N + 3)      # 격자 오른쪽 여분 열(원본 TC)


# ----------------------------------------------------------------------------- state → Params
def applicable(s, C) -> tuple:
    """(적용 가능 여부, 사유). 판단 근거는 상수·격자 사실뿐(임계값 없음)."""
    if C.TREE_FWD_RULE != "piecewise_quarter_step" or C.NODE_DISCOUNT_CONV != "1_discrete_fwd":
        return False, f"검토자 방식 아님(TREE_FWD_RULE={C.TREE_FWD_RULE}, NODE_DISCOUNT_CONV={C.NODE_DISCOUNT_CONV})"
    if C.INTERP_METHOD_PRE not in METHODS or (C.INTERP_METHOD, C.INTERP_SPACE_GRID) != ("linear", "log_df"):
        return False, f"보간 조합이 검토자 방식 아님({C.INTERP_METHOD_PRE}/{C.INTERP_METHOD}/{C.INTERP_SPACE_GRID})"
    step = s.provenance.grid_settings.step
    if step not in C.STEP_MODES:
        return False, f"격자 설정 없음(step={step!r})"
    P = round(1.0 / C.STEP_MODES[step])
    if P not in STEP_WORDS:
        return False, f"스텝 {step}(P={P})은 이표기간과 정수로 맞지 않음(주간·월간만)"
    N = int(s.grid.tree.N or 0)
    if N <= 0:
        return False, "격자 스텝 수 없음"
    for c in C.CURVE_IDS:
        m = C.RF_FREQ if c == "RF" else C.RD_FREQ
        if P % m or N % (P // m):
            return False, f"{c}: P={P}, m={m}, N={N} 이 정수로 맞지 않음"
        for k in s.grid.knots[c]:
            if abs(k["t"] * P - round(k["t"] * P)) > 1e-9:
                return False, f"{c}: 테너 {k['tenor']}(t={k['t']}) 가 스텝의 정수배가 아님"
    return True, "ok"


def _tenor_label(years: float) -> str:
    return f"{round(years * 12)} MTH" if years < 1 else f"{years:g} YEAR"


def _marker_label(years: float) -> str:
    return f"{round(years * 12)}mth" if years < 1 else f"{years:g}year"


def _qend_label(years: float) -> str:
    return f"{years:g}YR" if years < 1 else f"{years:g}year"


def params_from_state(s, C, c: str) -> Params:
    step = s.provenance.grid_settings.step
    P = round(1.0 / C.STEP_MODES[step])
    m = C.RF_FREQ if c == "RF" else C.RD_FREQ
    tenors = [{"label": _tenor_label(k["t"]), "years": k["t"], "step": int(round(k["t"] * P)), "ytm": k["ytm"]} for k in s.grid.knots[c]]
    cd = s.provenance.curve_date
    curve_date = datetime.fromisoformat(cd) if cd else None
    header = " · ".join(x for x in (s.provenance.valuation_date, s.run.curve_set_id, s.provenance.source_agency) if x)
    return Params(c, P, m, int(s.grid.tree.N), tenors, curve_date, header, C.PAR_FACE if hasattr(C, "PAR_FACE") else FACE, method=C.INTERP_METHOD_PRE)


def build(s, C, c: str) -> Params:
    ok, why = applicable(s, C)
    if not ok:
        raise ValueError(f"검토자 수식 시트 적용 불가: {why}")
    return params_from_state(s, C, c)


# ----------------------------------------------------------------------------- 라벨
def _adapt(text: str, p: Params) -> str:
    """원본 라벨의 WEEK/QUARTER 를 스텝·이표기간 단어로 치환(주간×분기면 원문 그대로)."""
    sw, pw = STEP_WORDS[p.P], PERIOD_WORDS.get(p.m, f"1/{p.m}Y")
    if sw != "WEEK":
        text = text.replace("WEEKLY", f"{sw}LY").replace("Weekly", f"{sw.title()}ly").replace("WEEKS", f"{sw}S").replace("Weekly", f"{sw.title()}ly")
    if pw != "QUARTER":
        text = text.replace("QUARTERLY", f"{pw}LY").replace("Quarterly", f"{pw.title()}ly").replace("QUARTER", pw)
    return text


def labels(p: Params) -> dict:
    """행 → B열 라벨(원본 문자열; 끝 공백·'FOMULA' 철자 그대로)."""
    texts = layout()["sheets"][p.sheet]["texts"]
    return {int(k[1:]): _adapt(v, p) for k, v in texts.items() if k.startswith("B") and k[1:].isdigit()}


# ----------------------------------------------------------------------------- 수식 생성 (셀 → 값 | '=수식')
def _interp_formula(p: Params, cl: str, r_idx: int) -> str:
    Lt, r6 = p.Lt, 6
    rng4, rng6 = f"$C$4:${Lt}$4", f"$C${r6}:${Lt}${r6}"
    mt = f"MATCH({cl}${r_idx},{rng4},1)"
    return (f"=IF({cl}${r_idx}<=$C$4,$C${r6},IF({cl}${r_idx}>=${Lt}$4,${Lt}${r6},"
            f"INDEX({rng6},{mt})+(INDEX({rng6},{mt}+1)-INDEX({rng6},{mt}))*({cl}${r_idx}-INDEX({rng4},{mt}))/(INDEX({rng4},{mt}+1)-INDEX({rng4},{mt}))))")


def _pchip_slope_formula(p: Params, i: int) -> str:
    """행 8: 마디 i(1-based)의 PCHIP 기울기 — reference/interp_ref.pchip_slopes 와 같은 식·같은 연산 순서(Fritsch–Carlson 조화평균, 끝점 3점식+형태 보존)."""
    n = p.nT
    X = lambda j: col(j + 2)
    h = lambda a, b: f"({X(b)}$4-{X(a)}$4)"
    mm = lambda a, b: f"(({X(b)}$6-{X(a)}$6)/({X(b)}$4-{X(a)}$4))"
    if i == 1 or i == n:
        if i == 1:
            h0, h1, m0, m1 = h(1, 2), h(2, 3), mm(1, 2), mm(2, 3)
        else:
            h0, h1, m0, m1 = h(n - 1, n), h(n - 2, n - 1), mm(n - 1, n), mm(n - 2, n - 1)
        d = f"(((2*{h0}+{h1})*{m0}-{h0}*{m1})/({h0}+{h1}))"
        return f"=IF(SIGN({d})<>SIGN({m0}),0,IF(AND(SIGN({m0})<>SIGN({m1}),ABS({d})>3*ABS({m0})),3*{m0},{d}))"
    hL, hR, mL, mR = h(i - 1, i), h(i, i + 1), mm(i - 1, i), mm(i, i + 1)
    w1, w2 = f"(2*{hR}+{hL})", f"({hR}+2*{hL})"
    return f"=IF(OR(SIGN({mL})<>SIGN({mR}),{mL}=0,{mR}=0),0,1/(({w1}/{mL}+{w2}/{mR})/({w1}+{w2})))"


def _hermite_formula(p: Params, cl: str) -> str:
    """행 11(PCHIP): 보조행 h_seg(구간 번호)·h_s(구간 내 위치 s)로 3차 Hermite — reference/interp_ref._Hermite._eval 과 같은 항 순서."""
    R = p.rows; Lt = p.Lt
    w, S, T = f"{cl}${R['idx']}", f"{cl}${R['h_seg']}", f"{cl}${R['h_s']}"
    x4, y6, d8 = f"$C$4:${Lt}$4", f"$C$6:${Lt}$6", f"$C$8:${Lt}$8"
    H = f"(INDEX({x4},{S}+1)-INDEX({x4},{S}))"
    return (f"=IF({w}<=$C$4,$C$6,IF({w}>=${Lt}$4,${Lt}$6,"
            f"(2*{T}^3-3*{T}^2+1)*INDEX({y6},{S})+({T}^3-2*{T}^2+{T})*{H}*INDEX({d8},{S})"
            f"+(-2*{T}^3+3*{T}^2)*INDEX({y6},{S}+1)+({T}^3-{T}^2)*{H}*INDEX({d8},{S}+1)))")


def cells(p: Params) -> dict:
    """{(row, col): 값 또는 '=수식'}. 열은 1-based(C=3)."""
    R, N, Nq, spq, P, m = p.rows, p.N, p.Nq, p.spq, p.P, p.m
    Lw, Lq, Xc, Lt = p.Lw, p.Lq, p.Xc, p.Lt
    out = {}
    lab = labels(p)
    for r, text in lab.items():
        out[(r, 2)] = text
    out[(1, 4)] = p.header
    out[(5, 2)] = p.curve_date
    face_ref = "C$21" if p.kind == "RF" else "$C$40"   # 원본: Rf 는 행 21(PV OF BOND 상수), Rd 는 C40(주간 부트스트랩 액면)
    pchip = p.method == "pchip"
    # 블록 1 테너
    for i, t in enumerate(p.tenors, start=1):
        c = i + 2; cl = col(c)
        out[(4, c)] = t["step"]; out[(5, c)] = t["label"]; out[(6, c)] = t["ytm"]
        out[(7, c)] = f"=INDEX($C${R['spot']}:${Lw}${R['spot']},{cl}$4)"
        out[(9, t["step"] + 2)] = _marker_label(t["years"])
        if pchip:
            out[(R["slope"], c)] = _pchip_slope_formula(p, i)
    if pchip:
        out[(R["slope"], 2)] = "PCHIP SLOPE (dY/dSTEP)"
        out[(R["h_seg"], 2)] = "PCHIP SEGMENT"; out[(R["h_s"], 2)] = "PCHIP s = (STEP-x_i)/h"
    # 블록 2 격자
    for w in range(1, N + 1):
        c = w + 2; cl = col(c); pv = col(c - 1)
        out[(R["idx"], c)] = 1 if w == 1 else f"={pv}{R['idx']}+1"
        if pchip:
            out[(R["h_seg"], c)] = f"=IFERROR(MATCH({cl}${R['idx']},$C$4:${Lt}$4,1),1)"
            out[(R["h_s"], c)] = f"=IF({cl}${R['h_seg']}>={p.nT},0,({cl}${R['idx']}-INDEX($C$4:${Lt}$4,{cl}${R['h_seg']}))/(INDEX($C$4:${Lt}$4,{cl}${R['h_seg']}+1)-INDEX($C$4:${Lt}$4,{cl}${R['h_seg']})))"
            out[(R["ytm"], c)] = _hermite_formula(p, cl)
        else:
            out[(R["ytm"], c)] = _interp_formula(p, cl, R["idx"])
        out[(R["spot"], c)] = f"=(1+{cl}{R['ws']})^{P}-1"
        q = -(-w // spq); Aq = col(spq * q + 2)
        if q == 1:
            out[(R["fwd"], c)] = f"=((1+${Aq}${R['ws']})^{spq})^(1/{spq})-1"
        else:
            Aq1 = col(spq * (q - 1) + 2)
            out[(R["fwd"], c)] = f"=((1+${Aq}${R['ws']})^{spq * q}/(1+${Aq1}${R['ws']})^{spq * (q - 1)})^(1/{spq})-1"
    # 블록 3 이표 격자 부트스트랩
    for q in range(1, Nq + 1):
        c = q + 2; cl = col(c); pv = col(c - 1)
        out[(R["q"], c)] = 1 if q == 1 else f"={pv}{R['q']}+1"
        out[(R["qw"], c)] = f"={cl}{R['q']}*{spq}"
        out[(R["qy"], c)] = f"=HLOOKUP({cl}{R['qw']},$C${R['idx']}:${Lw}${R['ytm']},2,FALSE)"
        if p.kind == "RF":
            out[(R["qc"], c)] = f"={cl}{R['qy']}/{m}"
            out[(R["qpp"], c)] = f"={cl}${R['qpb']}/(1+{cl}{R['qc']})^{cl}{R['q']}"
            out[(R["qpb"], c)] = p.face
            out[(R["qdf"], c)] = f"=({cl}{R['qpb']}-{cl}{R['qc']}*{cl}{R['qpb']}*{cl}{R['qs']})/({cl}{R['qpb']}+{cl}{R['qc']}*{cl}{R['qpb']})"
            out[(R["qmc"], c)] = f"={cl}{R['qc']}*{cl}{R['qpb']}*{cl}{R['qs']}+({cl}{R['qc']}*{cl}{R['qpb']}+{cl}{R['qpb']})*{cl}{R['qdf']}"
        else:
            cc = f"{cl}{R['qy']}/{m}"
            out[(R["qpb"], c)] = f"=PV({cc},{cl}{R['q']},-{cc}*{face_ref},-{face_ref})"
            out[(R["qdf"], c)] = f"=({cl}{R['qpb']}-{cc}*{cl}{R['qpb']}*{cl}{R['qs']})/({cl}{R['qpb']}+{cc}*{cl}{R['qpb']})"
            out[(R["qmc"], c)] = f"={cc}*{face_ref}*{cl}{R['qs']}+({cc}*{face_ref}+{face_ref})*{cl}{R['qdf']}"
        out[(R["qs"], c)] = 0 if q == 1 else f"={pv}{R['qs']}+{pv}{R['qdf']}"
        out[(R["qsp"], c)] = f"=(1/{cl}{R['qdf']})^(1/{cl}{R['q']})-1"
        out[(R["qsy"], c)] = f"={cl}{R['qy']}" if q == 1 else f"=(1+{cl}{R['qsp']})^{m}-1"
        out[(R["qfq"], c)] = f"={cl}{R['qsp']}" if q == 1 else f"={pv}{R['qdf']}/{cl}{R['qdf']}-1"
        out[(R["qfy"], c)] = f"={cl}{R['qsy']}" if q == 1 else f"=(1+{cl}{R['qfq']})^{m}-1"
        out[(R["mk"], spq * q + 2)] = _qend_label(q / m)
    # 블록 4 격자 위 선도·PVF·검증
    for w in range(1, N + 1):
        c = w + 2; cl = col(c); pv = col(c - 1); nx = col(c + 1)
        out[(R["widx"], c)] = 1 if w == 1 else f"={pv}{R['widx']}+1"
        if w % spq == 0:
            out[(R["ws"], c)] = f"=(1/{col(w // spq + 2)}{R['qdf']})^(1/{cl}{R['widx']})-1"
        else:
            out[(R["ws"], c)] = f"=({cl}{R['f2']}*{cl}{R['f1']})^(1/{cl}{R['widx']})-1"
        out[(R["f1"], c)] = f"=1+{cl}{R['fwd']}"
        out[(R["f2"], c)] = 1 if w == 1 else f"={pv}{R['f2']}*{pv}{R['f1']}"
        out[(R["wf"], c)] = f"={cl}{R['fwd']}"
        out[(R["pvf"], c)] = f"=1/{cl}{R['f1']}"
        if p.kind == "RD":
            out[(R["wfy"], c)] = f"=(1+{cl}{R['wf']})^{P}-1"
        if w == 1:
            out[(R["mc1"], c)] = f"=ROUND(SUMPRODUCT(ABS(D{R['mc1']}:{Lw}{R['mc1']})),9)=0"
            prod = f"{Xc}{R['pvf']}" if p.kind == "RF" else f"PRODUCT(C{R['pvf']}:{Lw}{R['pvf']})"
            out[(R["mc2"], c)] = f"=ROUND(D{R['mc2']}*C{R['pvf']}-{prod}*{face_ref.replace('C$21', '$C$21')},6)=0"
        else:
            out[(R["mc1"], c)] = f"={cl}{R['pvf']}*{cl}{R['f1']}-1"
            out[(R["mc2"], c)] = f"={nx}{R['mc2']}*{cl}{R['pvf']}"
    if p.kind == "RF":
        out[(R["ws"], N + 3)] = f"=1/(1+{Lw}{R['ws']})^{Lw}{R['widx']}"
        out[(R["pvf"], N + 3)] = f"=PRODUCT(C{R['pvf']}:{Lw}{R['pvf']})"
        out[(R["mc2"], N + 3)] = "=$C$21"   # 원본 TC38 은 #REF!(삭제된 참조) — 역방향 PV 의 출발값(액면)으로 복원
    else:
        out[(R["mc2"], N + 2)] = f"=$C$40*{Lw}{R['pvf']}"
        # 블록 5 주간(스텝) 부트스트랩 대조
        out[(R["face"], 3)] = p.face
        for w in range(1, N + 1):
            c = w + 2; cl = col(c); pv = col(c - 1)
            out[(R["widx2"], c)] = 1 if w == 1 else f"={pv}{R['widx2']}+1"
            out[(R["wc"], c)] = f"={cl}{R['ytm']}/{P}"
            out[(R["wdf"], c)] = f"=($C$40-{cl}{R['wc']}*$C$40*{cl}{R['wsum']})/($C$40+{cl}{R['wc']}*$C$40)"
            out[(R["wsum"], c)] = 0 if w == 1 else f"={pv}{R['wsum']}+{pv}{R['wdf']}"
            out[(R["wsp"], c)] = f"=(1/{cl}{R['wdf']})^(1/{cl}{R['widx2']})-1"
            out[(R["wmc"], c)] = f"={cl}{R['wc']}*$C$40*{cl}{R['wsum']}+({cl}{R['wc']}*$C$40+$C$40)*{cl}{R['wdf']}"
            out[(R["wf1"], c)] = f"=(1+{cl}{R['wsp']})^{cl}{R['widx2']}/{cl}{R['wf2']}"
            out[(R["wf2"], c)] = 1 if w == 1 else f"={pv}{R['wf2']}*{pv}{R['wf1']}"
            out[(R["wfw"], c)] = f"={cl}{R['wf1']}-1"
            out[(R["wpvf"], c)] = f"=1/{cl}{R['wf1']}"
            if w == 1:
                out[(R["wmc1"], c)] = f"=ROUND(SUMPRODUCT(ABS(D{R['wmc1']}:{Lw}{R['wmc1']})),9)=0"
                out[(R["wmc2"], c)] = f"=ROUND(SUMPRODUCT(ABS(D{R['wmc2']}:{Lw}{R['wmc2']})),6)=0"
            else:
                out[(R["wmc1"], c)] = f"={cl}{R['wpvf']}*{cl}{R['wf1']}-1"
                out[(R["wmc2"], c)] = f"={cl}{R['wmc']}-$C$40"
        out[(R["wpvf"], N + 3)] = f"=PRODUCT(C{R['wpvf']}:{Lw}{R['wpvf']})"
    return out


# ----------------------------------------------------------------------------- 파이썬 평가(같은 수식·같은 순서)
def evaluate(p: Params) -> dict:
    """행 번호 → 값 리스트(열 C 부터). 여분 열(TC) 값은 키 (row, 'X'). 수식 셀과 1:1 — 화면·검증용."""
    R, N, Nq, spq, P, m, F = p.rows, p.N, p.Nq, p.spq, p.P, p.m, p.face
    V = {}
    steps = [t["step"] for t in p.tenors]; ys = [t["ytm"] for t in p.tenors]
    V[4] = list(steps); V[5] = [t["label"] for t in p.tenors]; V[6] = list(ys)

    if p.method == "pchip":  # 시트와 같은 x(스텝 단위)로 참조 구현을 그대로 평가 — 엔진(연 단위)과는 부동소수 잡음만 다르다
        from ..reference import interp_ref as IR
        xs = [float(x) for x in steps]
        V[R["slope"]] = IR.pchip_slopes(xs, ys)
        fh = IR.pchip(xs, ys)

        def interp(w):
            if w <= steps[0]:
                return ys[0]
            if w >= steps[-1]:
                return ys[-1]
            return fh(float(w))
    else:
        def interp(w):
            if w <= steps[0]:
                return ys[0]
            if w >= steps[-1]:
                return ys[-1]
            i = max(j for j in range(len(steps)) if steps[j] <= w)
            return ys[i] + (ys[i + 1] - ys[i]) * (w - steps[i]) / (steps[i + 1] - steps[i])
    idx = list(range(1, N + 1)); ytm = [interp(w) for w in idx]
    V[R["idx"]] = idx; V[R["ytm"]] = ytm
    # 이표 격자
    qi = list(range(1, Nq + 1)); qw = [q * spq for q in qi]; qy = [ytm[w - 1] for w in qw]
    cpn = [y / m for y in qy]
    S, DF, Ssum = [], [], 0.0
    pb = []
    for q in qi:
        c = cpn[q - 1]
        if p.kind == "RF":
            P0 = F
        else:  # Excel PV(rate, nper, -pmt, -fv) = pmt·(1−(1+r)^−n)/r + fv·(1+r)^−n
            P0 = c * F * (1 - _xpow(1 + c, -q)) / c + F * _xpow(1 + c, -q)
        pb.append(P0)
        S.append(Ssum)
        d = (P0 - c * P0 * Ssum) / (P0 + c * P0)
        DF.append(d); Ssum = Ssum + d
    qsp = [(1 / DF[q - 1]) ** (1 / q) - 1 for q in qi]
    qsy = [qy[0]] + [_xpow(1 + qsp[q - 1], m) - 1 for q in qi[1:]]
    qfq = [qsp[0]] + [DF[q - 2] / DF[q - 1] - 1 for q in qi[1:]]
    qfy = [qsy[0]] + [_xpow(1 + qfq[q - 1], m) - 1 for q in qi[1:]]
    if p.kind == "RF":
        qmc = [c * F * s + (c * F + F) * d for c, s, d in zip(cpn, S, DF)]
        V[R["qc"]] = cpn; V[R["qpp"]] = [F / _xpow(1 + c, q) for c, q in zip(cpn, qi)]; V[R["qpb"]] = [F] * Nq
    else:
        qmc = [c * F * s + (c * F + F) * d for c, s, d in zip(cpn, S, DF)]
        V[R["qpb"]] = pb
    V[R["q"]] = qi; V[R["qw"]] = qw; V[R["qy"]] = qy; V[R["qdf"]] = DF; V[R["qs"]] = S; V[R["qsp"]] = qsp; V[R["qsy"]] = qsy; V[R["qfq"]] = qfq; V[R["qfy"]] = qfy; V[R["qmc"]] = qmc
    # 격자 선도(분기말 앵커) → 성장계수 → 스텝 현물 → 연현물
    anchor = {q: (1 / DF[q - 1]) ** (1 / (spq * q)) - 1 for q in qi}
    fwd = []
    for w in idx:
        q = -(-w // spq)
        num = _xpow(1 + anchor[q], spq * q)
        den = 1.0 if q == 1 else _xpow(1 + anchor[q - 1], spq * (q - 1))
        fwd.append((num / den) ** (1 / spq) - 1)
    f1 = [1 + f for f in fwd]
    f2, acc = [], 1.0
    for w in idx:
        f2.append(acc); acc = acc * f1[w - 1]
    ws = [anchor[w // spq] if w % spq == 0 else (f2[w - 1] * f1[w - 1]) ** (1 / w) - 1 for w in idx]
    spot = [_xpow(1 + r, P) - 1 for r in ws]
    pvf = [1 / x for x in f1]
    V[R["fwd"]] = fwd; V[R["spot"]] = spot; V[7] = [spot[t - 1] for t in steps]
    V[R["widx"]] = idx; V[R["ws"]] = ws; V[R["f1"]] = f1; V[R["f2"]] = f2; V[R["wf"]] = fwd; V[R["pvf"]] = pvf
    if p.kind == "RD":
        V[R["wfy"]] = [_xpow(1 + f, P) - 1 for f in fwd]
    mc1 = [None] + [v * g - 1 for v, g in zip(pvf[1:], f1[1:])]
    back = [0.0] * N; acc = F
    for i in range(N - 1, -1, -1):
        acc = acc * pvf[i]; back[i] = acc
    prod = math.prod(pvf)
    mc1[0] = round(math.fsum(abs(x) for x in mc1[1:]), 9) == 0
    mc2 = [round(back[1] * pvf[0] - prod * F, 6) == 0] + back[1:]
    V[R["mc1"]] = mc1; V[R["mc2"]] = mc2
    if p.kind == "RF":
        V[(R["ws"], "X")] = 1 / _xpow(1 + ws[-1], N); V[(R["pvf"], "X")] = prod; V[(R["mc2"], "X")] = F
    else:
        # 블록 5 스텝 부트스트랩
        wc = [y / P for y in ytm]
        wdf, wsum, Ssum = [], [], 0.0
        for w in idx:
            c = wc[w - 1]; wsum.append(Ssum)
            d = (F - c * F * Ssum) / (F + c * F); wdf.append(d); Ssum = Ssum + d
        wsp = [(1 / wdf[w - 1]) ** (1 / w) - 1 for w in idx]
        wmc = [c * F * s + (c * F + F) * d for c, s, d in zip(wc, wsum, wdf)]
        wf1, wf2, acc = [], [], 1.0
        for w in idx:
            wf2.append(acc); g = _xpow(1 + wsp[w - 1], w) / acc; wf1.append(g); acc = acc * g
        wfw = [g - 1 for g in wf1]; wpvf = [1 / g for g in wf1]
        wmc1 = [None] + [v * g - 1 for v, g in zip(wpvf[1:], wf1[1:])]; wmc1[0] = round(math.fsum(abs(x) for x in wmc1[1:]), 9) == 0
        wmc2 = [None] + [x - F for x in wmc[1:]]; wmc2[0] = round(math.fsum(abs(x) for x in wmc2[1:]), 6) == 0
        V[R["face"]] = [F]; V[R["widx2"]] = idx; V[R["wc"]] = wc; V[R["wdf"]] = wdf; V[R["wsum"]] = wsum; V[R["wsp"]] = wsp; V[R["wmc"]] = wmc
        V[R["wf1"]] = wf1; V[R["wf2"]] = wf2; V[R["wfw"]] = wfw; V[R["wpvf"]] = wpvf; V[R["wmc1"]] = wmc1; V[R["wmc2"]] = wmc2
        V[(R["wpvf"], "X")] = math.prod(wpvf)
    return V


# ----------------------------------------------------------------------------- 서식 대응
_W_CANDS = [(1, {"first", "q1"}), (2, {"q1"}), (13, {"q1", "tenor", "qend"}), (14, {"qstart"}), (15, set()), (26, {"tenor", "qend"}), (65, {"qend"}), (520, {"tenor", "qend", "last"})]
_Q_CANDS = [(1, {"first", "tenor"}), (2, {"tenor"}), (5, set()), (40, {"tenor", "last"})]
_T_CANDS = [(1, {"first"}), (5, set()), (12, {"last"})]
_WEIGHTS = {"first": 4, "last": 4, "q1": 2, "qstart": 2, "tenor": 1, "qend": 1}


def _closest(flags: set, cands) -> int:
    best, score = None, None
    for k, fs in cands:
        sc = sum(_WEIGHTS[f] for f in flags ^ fs)
        if score is None or sc < score:
            best, score = k, sc
    return best


def _orig_col(kind: str, c: int, p: Params):
    """대상 열(1-based) → 원본 열(1-based) 또는 None. 열의 성격이 같은 원본 열을 고른다."""
    if c <= 2:
        return c
    if kind == "fixed":
        return c if c <= 18 else None
    if kind == "tenor":
        i = c - 2
        if i <= p.nT:
            flags = {"first"} if i == 1 else ({"last"} if i == p.nT else set())
            return _closest(flags, _T_CANDS) + 2
        k = i - p.nT
        return 14 + k if 1 <= k <= 4 else None
    if kind == "quarterly":
        q = c - 2
        if 1 <= q <= p.Nq:
            flags = set()
            if q == 1: flags.add("first")
            if q == p.Nq: flags.add("last")
            if q * p.spq in p.tenor_steps: flags.add("tenor")
            return _closest(flags, _Q_CANDS) + 2
        return None
    w = c - 2
    if 1 <= w <= p.N:
        flags = set()
        if w == 1: flags.add("first")
        if w <= p.spq: flags.add("q1")
        if w in p.tenor_steps: flags.add("tenor")
        if w % p.spq == 0: flags.add("qend")
        if w > p.spq and w % p.spq == 1: flags.add("qstart")
        if w == p.N: flags.add("last")
        return _closest(flags, _W_CANDS) + 2
    if w == p.N + 1: return 523
    if w == p.N + 2: return 524
    return None


def _run_style(runs, c: int) -> int:
    for a, b, sid in runs:
        if a <= c <= b:
            return sid
    return -1


def _color(d):
    from openpyxl.styles import Color
    if not d:
        return None
    if "rgb" in d: return Color(rgb=d["rgb"])
    if "theme" in d: return Color(theme=d["theme"], tint=d["tint"])
    if "indexed" in d: return Color(indexed=d["indexed"])
    return None


def _style_objects(st: dict):
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    f = st["font"]; fl = st["fill"]; b = st["border"]; a = st["align"]
    font = Font(name=f[0], size=f[1], bold=f[2], italic=f[3], color=_color(f[4]))
    if fl[0]:
        kw = {"patternType": fl[0]}
        if fl[1]: kw["fgColor"] = _color(fl[1])
        if fl[2]: kw["bgColor"] = _color(fl[2])
        fill = PatternFill(**kw)
    else:
        fill = PatternFill()
    sides = [Side(style=x[0], color=_color(x[1])) if x else Side() for x in b]
    border = Border(left=sides[0], right=sides[1], top=sides[2], bottom=sides[3])
    align = Alignment(horizontal=a[0], vertical=a[1], indent=a[2], wrap_text=a[3] or None)
    return font, fill, border, align, st["nf"]


def number_format_of(p: Params, row: int) -> str:
    """행의 데이터 셀(원본 D열) 숫자 서식 — 화면 표시용."""
    L = layout()["sheets"][p.sheet]
    sid = _run_style(L["rows"][str(row)]["runs"], 4)
    return L["styles"][sid]["nf"] if sid >= 0 else "General"


# ----------------------------------------------------------------------------- 작성기
def write_sheet(ws, wb, p: Params) -> dict:
    """ws 에 수식·서식을 쓴다. 반환: cell_map 용 범위 {'RF:block1': 'B3:N7', 'RF:row22': 'C22:AP22', …}."""
    from openpyxl.styles import Color
    L = layout()["sheets"][p.sheet]
    styles = L["styles"]
    cache = {}

    def sty(sid):
        if sid not in cache:
            cache[sid] = _style_objects(styles[sid])
        return cache[sid]
    wb.loaded_theme = theme_bytes()
    ws.sheet_format.defaultColWidth = L["default_col_width"]; ws.sheet_format.defaultRowHeight = L["default_row_height"]; ws.sheet_format.customHeight = True
    ws.sheet_view.showGridLines = L["show_grid_lines"]; ws.sheet_view.zoomScaleNormal = L["zoom"]
    ws.sheet_properties.tabColor = _color(L["tab_color"])
    ws.freeze_panes = L["freeze_panes"]
    for k, w in L["col_widths"].items():
        ws.column_dimensions[k].width = w
    kinds = _KINDS[p.kind]
    max_col = p.N + 4
    data = cells(p)
    for r_s, spec in L["rows"].items():
        r = int(r_s)
        rd = ws.row_dimensions[r]
        rd.height = spec["h"]; rd.outlineLevel = spec["outline"]; rd.thickTop = spec["thickTop"] or None; rd.thickBot = spec["thickBot"] or None
        if spec["row_style"] is not None:
            font, fill, border, align, nf = sty(spec["row_style"])
            rd.font, rd.fill, rd.border, rd.alignment, rd.number_format = font, fill, border, align, nf
        kind = kinds.get(r, "fixed")
        runs = spec["runs"]
        for c in range(1, max_col + 1):
            oc = _orig_col(kind, c, p)
            sid = _run_style(runs, oc) if oc else -1
            v = data.get((r, c))
            if sid < 0 and v is None:
                continue
            cell = ws.cell(r, c)
            if sid >= 0:
                font, fill, border, align, nf = sty(sid)
                cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = font, fill, border, align, nf
            if v is not None:
                cell.value = v
    if p.method == "pchip":  # 원본 배치 밖의 행: 행 8 기울기(테너 행 6 서식), 보조행 2개(주 번호 행 10 서식) — 값·수식은 data 에 있다
        R = p.rows
        extra = [(R["slope"], "tenor", "6", "0.000E+00"), (R["h_seg"], "weekly", "10", "0"), (R["h_s"], "weekly", "10", "0.0000")]
        for r, kind, src_row, nf_data in extra:
            runs = L["rows"][src_row]["runs"]; lab_sid = _run_style(L["rows"]["11"]["runs"], 2)
            ws.row_dimensions[r].height = L["default_row_height"]
            for c in range(2, max_col + 1):
                oc = _orig_col(kind, c, p)
                sid = lab_sid if c == 2 else (_run_style(runs, oc) if oc else -1)
                v = data.get((r, c))
                if sid < 0 and v is None:
                    continue
                cell = ws.cell(r, c)
                if sid >= 0:
                    font, fill, border, align, nf = sty(sid)
                    cell.font, cell.fill, cell.border, cell.alignment = font, fill, border, align
                    cell.number_format = nf if c == 2 else nf_data
                if v is not None:
                    cell.value = v
    rng = {}
    for key, title_row, rows, kind in _BLOCKS[p.kind]:
        last = {"tenor": p.Lt, "quarterly": p.Lq, "weekly": p.Xc}[kind]
        rng[f"{p.kind}:{key}"] = f"B{title_row}:{last}{rows[-1]}"
        for r in rows:
            rng[f"{p.kind}:row{r}"] = f"C{r}:{last}{r}"
    return rng


def blocks(p: Params) -> list:
    """화면용 블록(검토자 배치 그대로: 행 = 라벨, 열 = 스텝). 값은 파이썬 모형(evaluate)."""
    V = evaluate(p); lab = labels(p); R = p.rows
    sw = STEP_WORDS[p.P].title()
    # 시트에서 라벨이 없는 행(주 번호·B5 고시일·Rd 행 34·액면)은 화면에서만 설명 라벨을 붙인다(xlsx 는 원본 그대로)
    screen_only = {5: (p.curve_date.date().isoformat() if p.curve_date else "curve date"), R["idx"]: f"{sw} no.", R["widx"]: f"{sw} no.", R["slope"]: "PCHIP SLOPE (dY/dSTEP)"}
    if p.kind == "RD":
        screen_only.update({R["widx2"]: f"{sw} no.", R["wfy"]: "FORWARD RATE - Yearly (grid, =(1+F)^P-1)", R["face"]: "FACE (PV OF BOND)"})
    out = []
    for bi, (key, title_row, rows, kind) in enumerate(_BLOCKS[p.kind], start=1):
        block = {"index": bi, "key": key, "title": lab.get(title_row, key), "orient": "rows", "rows": []}
        rows = list(rows) + ([R["slope"]] if bi == 1 and p.method == "pchip" else [])
        for r in rows:
            vals = V.get(r)
            if vals is None:
                continue
            block["rows"].append({"label": lab.get(r) or screen_only.get(r, ""), "key": f"row{r}", "fmt": "0.000E+00" if r == R["slope"] else number_format_of(p, r),
                                  "values": list(vals), "extra": V.get((r, "X"))})
        out.append(block)
    return out

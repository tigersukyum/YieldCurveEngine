# -*- coding: utf-8 -*-
"""
excel_recalc.py — 개발/검증 도구: openpyxl 로 쓴 통합문서(수식 포함)를 실제 Excel(COM) 로 열어 전체 재계산·저장한 뒤,
캐시된 계산값을 openpyxl(data_only=True) 로 읽어 돌려준다. 앱 실행 경로가 아니라 검증 전용(Excel 이 설치된 Windows 에서만).
사용: python cb_valuation/step1_curve/reference/excel_recalc.py <xlsx> [시트 …]   → 시트별 값 JSON 출력(요약)
"""
from __future__ import annotations
import json, os, subprocess, sys, tempfile

PS = r'''
$ErrorActionPreference = "Stop"
$x = New-Object -ComObject Excel.Application
$x.Visible = $false; $x.DisplayAlerts = $false
try {
  $wb = $x.Workbooks.Open("{PATH}")
  $x.CalculateFullRebuild()
  $wb.Save()
  $wb.Close($true)
  "OK"
} finally {
  $x.Quit(); [System.Runtime.Interopservices.Marshal]::ReleaseComObject($x) | Out-Null
}
'''


def recalc(path: str) -> None:
    """Excel 로 열어 전체 재계산 후 저장(캐시값 기록)."""
    script = PS.replace("{PATH}", os.path.abspath(path).replace("/", "\\"))
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as fh:
        fh.write(script); ps1 = fh.name
    try:
        out = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", ps1], capture_output=True, text=True, timeout=600)
        if out.returncode != 0 or "OK" not in out.stdout:
            raise RuntimeError(f"Excel 재계산 실패: {out.stdout[-500:]} {out.stderr[-500:]}")
    finally:
        os.unlink(ps1)


def read_values(path: str, sheets=None) -> dict:
    """재계산 후 캐시값 읽기: {sheet: {cell: value}}. 수식 셀은 계산값, 상수 셀은 그대로."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = {}
    for name in (sheets or wb.sheetnames):
        ws = wb[name]
        out[name] = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    return out


def recalc_and_read(path: str, sheets=None) -> dict:
    recalc(path)
    return read_values(path, sheets)


if __name__ == "__main__":
    p = sys.argv[1]; sheets = sys.argv[2:] or None
    vals = recalc_and_read(p, sheets)
    print(json.dumps({k: {"cells": len(v), "sample": dict(list(v.items())[:5])} for k, v in vals.items()}, ensure_ascii=False, indent=1, default=str))

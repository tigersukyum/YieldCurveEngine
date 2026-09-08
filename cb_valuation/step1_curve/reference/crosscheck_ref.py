# -*- coding: utf-8 -*-
"""
crosscheck_ref.py — 독립 검증 도구(ref/ 원본이 있는 PC 에서만): 값복사된 참조 커브들과 엔진 부트스트랩을 대조한다.
  A. 평가사(KIS-NET) 공시 현물 커브 3종('연복리 연속 Spot'·'분기복리 이산 Spot'·'연 복리 이산 Spot', 영풍 FY25 패키지)  vs  엔진(같은 YTM Matrix 로 부트스트랩)
     — 국고채 + 회사채 전 등급(공모·사모), 테너별 차이(bp). 엔진 프로필: LINEAR(RF 반기·RD 분기), REVIEWER_2024(RF·RD 분기), PCHIP.
  B. KBI 2024 검토자 Rf_dc/Rd_dc(분기·선형)  vs  엔진 REVIEWER_2024 — DF·연현물·주간 선도·주간 현물.
  C1. 영풍 FY25 'Boot Strapping'(국고채 주간 4Y, 연속 PCHIP mode-B + 분기 별도 부트스트랩)  vs  엔진 LINEAR/PCHIP RF — 마디·격자.
  C2. KBI xlsm BOOT(값복사 마디, RF 반기·RD 분기 선형)  vs  엔진 LINEAR(같은 마디 합성 매트릭스) — 이표격자 YTM·기간현물·DF·0.25 격자.
  C3. KBI xlsm 'Bootstrapping_평가자'(YTM 복리변환, 주간 선형)  vs  엔진 LINEAR — 부트스트랩이 아님을 수치로 확인.
  D.  대교 'Par검증' 템플릿 예제 — 관례별 par 항등식.
사용: python cb_valuation/step1_curve/reference/crosscheck_ref.py [A|B|C|C1|C2|C3|D|all] → 표를 stdout, JSON 을 exports/crosscheck_ref.json
"""
from __future__ import annotations
import json, math, os, sys, tempfile, shutil

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
from cb_valuation.step1_curve.graph import step1_graph as G  # noqa: E402
from cb_valuation.step1_curve.app import runner as R  # noqa: E402
from cb_valuation.step1_curve.io.matrix_parser import preview, parse_matrix_text  # noqa: E402

REF = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "ref")
YP = os.path.join(REF, "영풍 주주간 계약 옵션", "8521_와이피씨_FY25_8521_평가보고서 검토자의 검토요구사항 (2025) (M)(서명)_20260311.xlsx")
KBI24 = os.path.join(REF, "KBI메탈 전환사채", "8521_평가보고서 검토자의 검토요구사항_Call Option Valuation_KBI metal_2024.xlsx")
TEN = G.Constants.TENOR_LABELS
BP = 1e4


def run_engine(matrix_text, profile, rf_row, rd_row, valuation, curve_date=None, step="weekly", horizon=10):
    base = tempfile.mkdtemp(prefix="cb_xc_")
    try:
        form = {"profile": profile, "matrix_text": matrix_text, "valuation_date": valuation, "curve_date": curve_date or valuation, "curve_set_id": "XC",
                "operator": "crosscheck", "source_agency": "KIS", "downloaded_at": valuation + "T09:00:00+09:00", "step": step, "horizon_years": horizon,
                "rf_row_index": rf_row, "rd_row_index": rd_row}
        s, C = R.prepare_state(form, base)
        s, _ = R.advance(s, C, base)
        while s.run.status == "paused":
            kind = s.run.paused_at_node.replace("approve_", "")
            ack = [f["code"] for f in s[f"approval_{kind}"].flags_seen] if kind == "exception" else []
            s, _ = R.decide_and_resume(s, C, base, "approved", "crosscheck", "auto", ack)
        if s.run.status != "done":
            raise RuntimeError(f"엔진 실패 {s.result.fail_code}: {s.result.fail_reason}")
        return s, C
    finally:
        shutil.rmtree(base, ignore_errors=True)


def sheet_rows(ws):
    """(구분, 적용대상채권) 라벨 → {tenor: value} ; 헤더 1행, 연수 2행, 값 3행부터(단위 %)."""
    out = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r is None or len(r) < 4 or (r[1] is None and r[2] is None):
            continue
        key = str(r[2] or "").strip()  # 적용대상채권 라벨은 시트 안에서 유일(사모 회사채는 '사모회사채AAA' 식으로 구분됨)
        out[key] = {t: (float(v) if isinstance(v, (int, float)) else None) for t, v in zip(TEN, r[3:3 + len(TEN)])}
    return out


def kis_label(row: dict) -> str:
    """엔진 파서 행 → KIS 시트 라벨: 사모무보증 블록의 '회사채AAA' 는 KIS 시트에서 '사모회사채AAA'."""
    lab = (row.get("label") or "").strip()
    return ("사모" + lab) if (row.get("block") == "사모무보증" and not lab.startswith("사모")) else lab


def matrix_text_from_sheet(ws):
    lines = []
    for i, r in enumerate(ws.iter_rows(min_row=1, values_only=True), 1):
        if i == 2 and all(isinstance(v, (int, float)) or v is None for v in r):
            continue  # 연수 행(fixture 에는 없음)
        cells = ["" if v is None else (str(v)) for v in r[:19]]
        lines.append(",".join(cells))
    return "\n".join(lines) + "\n"


# ----------------------------------------------------------------------------- A. KIS-NET 공시 현물 vs 엔진
def part_A(profiles=("LINEAR", "REVIEWER_2024", "PCHIP"), max_rd_rows=None):
    from openpyxl import load_workbook
    wb = load_workbook(YP, read_only=True, data_only=True)
    # 'YTM Matrix' 시트는 fixture A(2025-12-31 KIS-NET)와 같은 매트릭스(일부 행은 라벨 칸이 한 칸 밀려 파서가 못 읽음) → 값 동일성을 확인한 뒤 fixture A 원문을 입력으로 쓴다
    fixture = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tests", "fixtures", "kisnet_matrix_20251231.csv")
    with open(fixture, encoding="utf-8-sig") as fh:
        text = fh.read()
    pv = preview(text, G.Constants)
    parsed = {r["row_index"]: r for r in parse_matrix_text(text, G.Constants)["rows"]}  # ytm_pct 는 파서 원문에만 있다
    sheet_vals = sheet_rows(wb["YTM Matrix"]); same = diff = 0
    for r in pv["rows"]:
        key = kis_label(r)
        cand = [k for k in sheet_vals if k == key]
        if not cand:
            continue
        for t in TEN:
            a, b = parsed[r["row_index"]]["ytm_pct"].get(t), sheet_vals[cand[0]].get(t)
            if a is None and b is None:
                continue
            if a is not None and b is not None and abs(a - b) < 1e-9:
                same += 1
            else:
                diff += 1
    print(f"  fixture A vs 'YTM Matrix' 시트 값 대조: 일치 {same}, 불일치 {diff}")
    kis = {"cont": sheet_rows(wb["연복리 연속 Spot"]), "q": sheet_rows(wb["분기복리 이산 Spot"]), "ann": sheet_rows(wb["연 복리 이산 Spot"])}
    rows_by_idx = {r["row_index"]: r for r in pv["rows"]}
    rf_idx = pv["rf_candidates"][0]
    rd_rows = pv["rd_candidates"] if max_rd_rows is None else pv["rd_candidates"][:max_rd_rows]
    results = {}
    print(f"[A] 영풍 FY25 'YTM Matrix'({pv['n_rows']}행) → 엔진 부트스트랩 vs KIS-NET 공시 현물(연속/분기/연복리). 단위 bp(=0.01%p). 테너 = 엔진 마디(이표 격자 위)")
    for profile in profiles:
        C = G.Constants.with_profile(profile)
        print(f"\n--- 프로필 {profile}: RF_FREQ {C.RF_FREQ}·RD_FREQ {C.RD_FREQ}, 마디 보간 {C.INTERP_METHOD_PRE}")
        table = []
        cache = {}
        for rd_idx in rd_rows:
            s, Cc = run_engine(text, profile, rf_idx, rd_idx, "2025-12-31")
            for c in ("RF", "RD"):
                row = rows_by_idx[rf_idx if c == "RF" else rd_idx]
                key = kis_label(row)
                cand = [k for k in kis["cont"] if k == key]  # KIS 시트 라벨 매칭(적용대상채권; 사모는 '사모' 접두)
                if c == "RF" and profile in cache:
                    continue
                if not cand:
                    table.append({"curve": c, "row": row["row_index"], "label": key, "note": "KIS 시트에 같은 라벨 없음"}); continue
                kk = cand[0]
                m = Cc.RF_FREQ if c == "RF" else Cc.RD_FREQ
                ann = dict(zip([round(t, 8) for t in s.conv.spot_annual[c]["times"]], s.conv.spot_annual[c]["values"]))
                cont = dict(zip([round(t, 8) for t in s.conv.spot_cont[c]["times"]], s.conv.spot_cont[c]["values"]))
                diffs = {"cont": [], "q": [], "ann": []}
                for k in s.grid.knots[c]:
                    t = round(k["t"], 8); lab = k["tenor"]
                    if t not in ann:
                        continue
                    za, zc = ann[t], cont[t]
                    zq = 4 * ((1 + za) ** 0.25 - 1)  # 분기복리 명목 연율
                    for kind, ours in (("cont", zc), ("q", zq), ("ann", za)):
                        ref = kis[kind][kk].get(lab)
                        if ref is not None:
                            diffs[kind].append((lab, (ours * 100 - ref) * 100))  # bp
                rec = {"curve": c, "row": row["row_index"], "label": kk, "n": len(diffs["ann"])}
                for kind in ("cont", "q", "ann"):
                    d = [abs(x[1]) for x in diffs[kind]]
                    rec[f"max_{kind}"] = max(d) if d else None; rec[f"mean_{kind}"] = (sum(d) / len(d)) if d else None
                    rec[f"detail_{kind}"] = {lab: round(v, 3) for lab, v in diffs[kind]}
                table.append(rec)
                if c == "RF":
                    cache[profile] = True
        results[profile] = table
        print(f"{'커브':4} {'행':>3} {'라벨':22} {'n':>2} | 연속 max/mean | 분기 max/mean | 연복리 max/mean (bp)")
        for rec in table:
            if "note" in rec:
                print(f"{rec['curve']:4} {rec['row']:>3} {str(rec['label'])[:22]:22} {rec['note']}"); continue
            f = lambda k: f"{rec['max_' + k]:5.2f}/{rec['mean_' + k]:5.2f}" if rec['max_' + k] is not None else "  -  "
            print(f"{rec['curve']:4} {rec['row']:>3} {str(rec['label'])[:22]:22} {rec['n']:>2} | {f('cont'):>12} | {f('q'):>12} | {f('ann'):>12}")
        worst = max((rec["max_ann"] for rec in table if "max_ann" in rec and rec["max_ann"] is not None), default=None)
        print(f"  → 연복리 현물 |차이| 최대 {worst:.2f} bp" if worst is not None else "  → 없음")
        rf = next((rec for rec in table if rec["curve"] == "RF" and "detail_ann" in rec), None)
        if rf:
            print("  RF 테너별 연복리 차이(bp):", rf["detail_ann"])
            print("  RF 테너별 분기복리 차이(bp):", rf["detail_q"])
    wb.close()
    return results


# ----------------------------------------------------------------------------- B. KBI 2024 검토자 Rf_dc/Rd_dc vs 엔진 REVIEWER_2024
def part_B():
    from openpyxl import load_workbook
    wb = load_workbook(KBI24, read_only=True, data_only=True)
    rf = list(wb["Rf_dc"].iter_rows(min_row=1, max_row=38, max_col=524, values_only=True))
    rd = list(wb["Rd_dc"].iter_rows(min_row=1, max_row=53, max_col=524, values_only=True))
    wb.close()
    ten12 = TEN[:12]
    rf_ytm = [rf[5][c] for c in range(2, 14)]; rd_ytm = [rd[5][c] for c in range(2, 14)]
    # 합성 매트릭스(두 행): 검토자 시트 행 6 의 YTM(소수) → %
    lines = ["종류,구분,적용대상채권," + ",".join(TEN),
             ",국고채권,국고채," + ",".join(f"{v * 100:.6f}" for v in rf_ytm) + ",-,-,-,-",
             "회사채,사모무보증,회사채BB+," + ",".join(f"{v * 100:.6f}" for v in rd_ytm) + ",-,-,-,-"]
    text = "\n".join(lines) + "\n"
    pv = preview(text, G.Constants)
    s, C = run_engine(text, "REVIEWER_2024", pv["rf_candidates"][0], pv["rd_candidates"][0], "2024-12-31")
    out = {}
    print("\n[B] KBI 2024 검토자 Rf_dc/Rd_dc(값복사) vs 엔진 REVIEWER_2024(RF·RD 분기, 선형) — 전 열 대조")
    for c, sh, rows in (("RF", rf, dict(df=22, sy=25, fwd=13, spot=12, qy=18)), ("RD", rd, dict(df=20, sy=23, fwd=13, spot=12, qy=18))):
        df_sheet = [sh[rows["df"] - 1][j] for j in range(2, 42)]
        d_df = max(abs(a - b) for a, b in zip(s.bootstrap.df[c], df_sheet))
        qy_sheet = [sh[rows["qy"] - 1][j] for j in range(2, 42)]
        d_qy = max(abs(a - b) for a, b in zip(s.interp.ytm_on_coupon_grid[c]["values"], qy_sheet))
        sy_sheet = [sh[rows["sy"] - 1][j] for j in range(3, 42)]  # C 열은 원본이 YTM 앵커 → 제외
        d_sy = max(abs(a - b) for a, b in zip(s.conv.spot_annual[c]["values"][1:], sy_sheet))
        fwd_sheet = [sh[rows["fwd"] - 1][j] for j in range(2, 522)]
        d_fwd = max(abs(a - b) for a, b in zip(s.fwd.disc_per_step[c]["values"], fwd_sheet))
        # 주간 현물(연복리): 시트 행 12 는 (1+r_w)^52-1 체인; 엔진 tree.spot_annual_on_grid 는 log-DF 선형 → 정의가 같은 분기말에서만 비교
        spot_sheet = [sh[rows["spot"] - 1][j] for j in range(2, 522)]
        grid = s.tree.spot_annual_on_grid[c]["values"][1:]
        d_spot_q = max(abs(grid[13 * q - 1] - spot_sheet[13 * q - 1]) for q in range(2, 41))
        out[c] = dict(df=d_df, ytm_grid=d_qy, spot_yearly=d_sy, fwd_weekly=d_fwd, spot_weekly_qend=d_spot_q)
        print(f"  {c}: DF max|Δ| {d_df:.2e} | 이표격자 YTM {d_qy:.2e} | 연현물(D~AP) {d_sy:.2e} | 주간 선도 {d_fwd:.2e} | 분기말 주간 연현물 {d_spot_q:.2e}")
    return out


# ----------------------------------------------------------------------------- 공통(C·D)
FIXA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tests", "fixtures", "kisnet_matrix_20251231.csv")
KBI_XLSM = os.path.join(REF, "KBI메탈 전환사채", "감사인 질의 및 회신", "KBI기말_Call_감사인제출_0204 - 복사본.xlsm")
DAEKYO = os.path.join(REF, "대교_질문지 예시", "(DVC)Derivatives valuation request form_주주간계약_대교_251229_yang(수정) (1).xlsx")
TEN_T = {"3M": 0.25, "6M": 0.5, "9M": 0.75, "1Y": 1, "1.5Y": 1.5, "2Y": 2, "2.5Y": 2.5, "3Y": 3, "4Y": 4, "5Y": 5, "7Y": 7, "10Y": 10, "15Y": 15, "20Y": 20, "30Y": 30, "50Y": 50}


def fixture_text():
    with open(FIXA, encoding="utf-8-sig") as fh:
        return fh.read()


def row_index_by_label(text, label):
    for r in preview(text, G.Constants)["rows"]:
        if kis_label(r) == label:
            return r["row_index"]
    raise KeyError(label)


def _vals(x):
    return x["values"] if isinstance(x, dict) and "values" in x else list(x)


def _stats(d):
    d = [x for x in d if x is not None and math.isfinite(x)]
    if not d:
        return {"n": 0, "max": None, "mean": None}
    a = [abs(x) for x in d]
    return {"n": len(d), "max": max(a), "mean": sum(a) / len(a), "argmax": max(range(len(a)), key=lambda i: a[i])}


def _fmt(st, unit=""):
    return f"max {st['max']:.3g}{unit} mean {st['mean']:.3g}{unit} (n={st['n']})" if st["n"] else "n=0"


def _sheet_col(ws, rng):
    return [c.value for row in ws[rng] for c in row]


# ----------------------------------------------------------------------------- C1. 영풍 FY25 'Boot Strapping'(국고채·주간·4Y) vs 엔진 RF
def part_C1(profiles=("LINEAR", "PCHIP")):
    """시트: A 노드, B t=round(k/52,12), C 연속현물(PCHIP mode-B, 앵커 0.02165, 0.75Y 특이), E DF, H 분기복리 현물(별도 부트스트랩), J DF(Q), R 선형 YTM."""
    from openpyxl import load_workbook
    wb = load_workbook(YP, read_only=True, data_only=True)
    rows = list(wb["Boot Strapping"].iter_rows(min_row=3, max_row=211, max_col=21, values_only=True))
    wb.close()
    sh = {"t": [r[1] for r in rows], "C": [r[2] for r in rows], "E": [r[4] for r in rows], "H": [r[7] for r in rows], "J": [r[9] for r in rows], "R": [r[17] for r in rows]}
    text = fixture_text(); pv = preview(text, G.Constants)
    KN = [0.25, 0.5, 0.75, 1, 1.5, 2, 2.5, 3, 4]
    knot_nodes = {round(t * 52) for t in KN}
    out = {}
    print("\n[C1] 영풍 FY25 'Boot Strapping'(국고채, 주간 1/52, 4Y, 값복사) vs 엔진 RF(같은 YTM Matrix = fixture A). 시트 C 커브 = 연속현물 PCHIP(mode-B, t=0 앵커 0.02165, 0.75Y '현물=YTM'), Q 커브 = 분기복리 별도 부트스트랩(확정마디 선형·미지수 평탄)")
    for prof in profiles:
        s, C = run_engine(text, prof, pv["rf_candidates"][0], pv["rd_candidates"][0], "2025-12-31")
        sc = _vals(s.tree.spot_cont_on_grid["RF"]); df = _vals(s.tree.df_spot_on_grid["RF"]); yg = _vals(s.tree.ytm_on_grid["RF"])
        tg = s.grid.tree.times
        assert all(abs(tg[k] - sh["t"][k]) < 1e-6 for k in range(209)), "격자 불일치"  # 엔진 t 는 T_ROUND_DIGITS 자리 반올림
        rec = {"knots": [], "grid": {}}
        for t in KN:
            k = round(t * 52)
            zc = sc[k]; zq = 4 * (math.exp(zc / 4) - 1)
            rec["knots"].append({"t": t, "engine_cont": zc, "sheet_C": sh["C"][k], "d_C_bp": (zc - sh["C"][k]) * BP,
                                 "engine_q": zq, "sheet_H": sh["H"][k], "d_H_bp": (zq - sh["H"][k]) * BP,
                                 "engine_df": df[k], "d_E": df[k] - sh["E"][k], "d_J": df[k] - sh["J"][k]})
        eq = [4 * (math.exp(v / 4) - 1) for v in sc]
        for name, eng, ref, scale in (("C_cont_bp", sc, sh["C"], BP), ("H_q_bp", eq, sh["H"], BP), ("E_df", df, sh["E"], 1.0), ("J_df", df, sh["J"], 1.0), ("R_ytm_bp", yg, sh["R"], BP)):
            rec["grid"][name] = {
                "all_1_208": _stats([(eng[k] - ref[k]) * scale for k in range(1, 209)]),
                "t_lt_0.25": _stats([(eng[k] - ref[k]) * scale for k in range(1, 13)]),
                "t_ge_0.25_between_knots": _stats([(eng[k] - ref[k]) * scale for k in range(13, 209) if k not in knot_nodes]),
                "knots": _stats([(eng[k] - ref[k]) * scale for k in sorted(knot_nodes)]),
            }
        out[prof] = rec
        print(f"\n--- 프로필 {prof} (RF_FREQ {C.RF_FREQ}, 마디 보간 {C.INTERP_METHOD_PRE}, 격자 공간 {C.INTERP_SPACE_GRID if hasattr(C, 'INTERP_SPACE_GRID') else s.interp.space_grid})")
        print(f"{'t':>5} | {'엔진 연속':>10} {'시트 C':>10} {'Δbp':>7} | {'엔진 분기':>10} {'시트 H':>10} {'Δbp':>7} | {'ΔDF(E)':>9} {'ΔDF(J)':>9}")
        for r in rec["knots"]:
            print(f"{r['t']:5} | {r['engine_cont']:10.6f} {r['sheet_C']:10.6f} {r['d_C_bp']:+7.3f} | {r['engine_q']:10.6f} {r['sheet_H']:10.6f} {r['d_H_bp']:+7.3f} | {r['d_E']:+9.2e} {r['d_J']:+9.2e}")
        for name in ("C_cont_bp", "H_q_bp", "E_df", "R_ytm_bp"):
            g = rec["grid"][name]
            print(f"  격자 {name:10}: 전체 {_fmt(g['all_1_208'])} | t<0.25 {_fmt(g['t_lt_0.25'])} | 마디 {_fmt(g['knots'])} | 마디 사이(t≥0.25) {_fmt(g['t_ge_0.25_between_knots'])}")
    return out


# ----------------------------------------------------------------------------- C2. KBI xlsm BOOT(값복사 마디 G/V, 2024-12-31 잔존) vs 엔진 LINEAR
def part_C2():
    """BOOT: RF 반기 격자 E(0.5..10)·G=YTM/2 선형보간(값)·H 기간현물·I DF ; RD 분기 격자 T·V=YTM/4·W·X ; 0.25 격자 K/M/N(RF 연복리·연속), Z/AB/AC(RD).
    마디 YTM(G×2, V×4)로 합성 매트릭스를 만들어 엔진 LINEAR(RF 반기·RD 분기·선형) 를 돌린다. 3M RF 는 시트 C10(라이브 2.422%) 을 그대로 쓴다."""
    from openpyxl import load_workbook
    wb = load_workbook(KBI_XLSM, read_only=True, data_only=True)
    ws = wb["BOOT"]
    E, Gc, H, I = (_sheet_col(ws, r) for r in ("E10:E29", "G10:G29", "H10:H29", "I10:I29"))
    K, L, M, N = (_sheet_col(ws, r) for r in ("K10:K49", "L10:L49", "M10:M49", "N10:N49"))
    T, V, W, X = (_sheet_col(ws, r) for r in ("T10:T49", "V10:V49", "W10:W49", "X10:X49"))
    Z, AA, AB, AC = (_sheet_col(ws, r) for r in ("Z10:Z49", "AA10:AA49", "AB10:AB49", "AC10:AC49"))
    live3m = _sheet_col(ws, "C10:C10")[0]
    wb.close()
    rf = {"3M": live3m * 100}
    for lab, t in TEN_T.items():
        if t in (0.5, 1, 1.5, 2, 2.5, 3, 4, 5, 7, 10):
            rf[lab] = Gc[int(round(t * 2)) - 1] * 2 * 100
    rd = {lab: V[int(round(t * 4)) - 1] * 4 * 100 for lab, t in TEN_T.items() if t <= 10}
    lines = ["종류,구분,적용대상채권," + ",".join(TEN),
             ",국고채권,국고채," + ",".join(f"{rf[t]:.6f}" if t in rf else "-" for t in TEN),
             "회사채,사모무보증,회사채BB+," + ",".join(f"{rd[t]:.6f}" if t in rd else "-" for t in TEN)]
    text = "\n".join(lines) + "\n"
    pv = preview(text, G.Constants)
    s, C = run_engine(text, "LINEAR", pv["rf_candidates"][0], pv["rd_candidates"][0], "2024-12-31")
    out = {"inputs": {"rf_pct": rf, "rd_pct": rd}}
    print("\n[C2] KBI xlsm BOOT(값복사 마디 G/V = 2024-12-31 잔존값; RF 반기·RD 분기·YTM/m 선형) vs 엔진 LINEAR(같은 마디 YTM 합성 매트릭스)")
    for c, m, grid_t, ytm_m, spot_pp, dfs, q_t, q_pp, q_ann, q_cont in (("RF", 2, E, Gc, H, I, K, L, M, N), ("RD", 4, T, V, W, X, Z, AA, AB, AC)):
        eng_y = _vals(s.interp.ytm_on_coupon_grid[c]); eng_pp = _vals(s.bootstrap.spot_pp[c]); eng_df = _vals(s.bootstrap.df[c])
        n = len(grid_t)
        d_y = _stats([eng_y[i] - ytm_m[i] * m for i in range(n)])
        d_pp = _stats([eng_pp[i] - spot_pp[i] for i in range(n)])
        d_df = _stats([eng_df[i] - dfs[i] for i in range(n)])
        sa = _vals(s.tree.spot_annual_on_grid[c]); sco = _vals(s.tree.spot_cont_on_grid[c])
        idx = [int(round(t * 52)) for t in q_t]
        on_c = [j for j, t in enumerate(q_t) if abs(t * m - round(t * m)) < 1e-9]  # 이표 격자 위 점
        mid = [j for j in range(len(q_t)) if j not in on_c]
        d_ann_c = _stats([(sa[idx[j]] - q_ann[j]) * BP for j in on_c]); d_ann_m = _stats([(sa[idx[j]] - q_ann[j]) * BP for j in mid])
        d_cont_c = _stats([(sco[idx[j]] - q_cont[j]) * BP for j in on_c]); d_cont_m = _stats([(sco[idx[j]] - q_cont[j]) * BP for j in mid])
        out[c] = {"ytm_on_coupon_grid": d_y, "spot_per_period": d_pp, "df": d_df, "quarter_grid_annual_bp_on_coupon": d_ann_c, "quarter_grid_annual_bp_midpoints": d_ann_m,
                  "quarter_grid_cont_bp_on_coupon": d_cont_c, "quarter_grid_cont_bp_midpoints": d_cont_m,
                  "midpoint_detail_bp": {str(q_t[j]): round((sa[idx[j]] - q_ann[j]) * BP, 4) for j in mid}}
        print(f"  {c} (m={m}, {n}기): 이표격자 YTM |Δ| {_fmt(d_y)} | 기간현물 |Δ| {_fmt(d_pp)} | DF |Δ| {_fmt(d_df)}")
        print(f"     0.25 격자 연복리(M/AB) 이표점 {_fmt(d_ann_c, 'bp')} · 중점 {_fmt(d_ann_m, 'bp')} | 연속(N/AC) 이표점 {_fmt(d_cont_c, 'bp')} · 중점 {_fmt(d_cont_m, 'bp')}")
    return out


# ----------------------------------------------------------------------------- C3. KBI xlsm 'Bootstrapping_평가자'(복사본; 라이브 2025-12-31 row2/row58) vs 엔진 LINEAR
def part_C3():
    """시트: 마디 주(C74:N74) · 연속현물(C82:N82 RF, C203:N203 RD) · 주간 격자 연속현물(B102:TB102 RF, B223:TB223 RD; 0~520주).
    시트의 '현물' 은 부트스트랩이 아니라 YTM 의 복리 변환((1+y/m)^m−1; 3M·9M 은 단리 스텁) → 엔진(진짜 par 부트스트랩)과의 차이는 '이표 재투자 효과' 그 자체."""
    from openpyxl import load_workbook
    wb = load_workbook(KBI_XLSM, read_only=True, data_only=True)
    ws = wb["Bootstrapping_평가자"]
    weeks = [int(w) for w in _sheet_col(ws, "C74:N74")]
    rf_y, rf_knot, rf_week = _sheet_col(ws, "C76:N76"), _sheet_col(ws, "C82:N82"), _sheet_col(ws, "B102:TB102")
    rd_y, rd_knot, rd_week = _sheet_col(ws, "C197:N197"), _sheet_col(ws, "C203:N203"), _sheet_col(ws, "B223:TB223")
    wb.close()
    text = fixture_text(); pv = preview(text, G.Constants)
    rd_idx = row_index_by_label(text, "사모회사채BB+")
    s, C = run_engine(text, "LINEAR", pv["rf_candidates"][0], rd_idx, "2025-12-31")
    out = {}
    print("\n[C3] KBI xlsm 'Bootstrapping_평가자'(값·수식 혼재, 라이브 KIS-NET row2 국고채 / row58 사모 BB+) vs 엔진 LINEAR. 시트 현물 = YTM 복리변환(부트스트랩 아님)")
    rd_week_is_rf = all((a == b) for a, b in zip(rf_week, rd_week))
    out["rd_weekly_block_equals_rf"] = rd_week_is_rf
    if rd_week_is_rf:
        print("  주의: RD 주간 블록(행 221~223)이 RF 주간 블록(행 100~102)과 521셀 모두 동일 → 시트의 참조 버그(RD 주간 격자 = RF 값). RD 주간 비교는 생략")
    for c, y, knot, week, m in (("RF", rf_y, rf_knot, rf_week, 2), ("RD", rd_y, rd_knot, (None if rd_week_is_rf else rd_week), 4)):
        sc = _vals(s.tree.spot_cont_on_grid[c])
        # (1) 시트 마디 현물이 YTM 변환인지 재확인
        conv = []
        for i, w in enumerate(weeks):
            t = w / 52
            if m == 2 and abs(t - 0.25) < 1e-9:
                conv.append(4 * math.log1p(y[i] / 4))
            elif m == 2 and abs(t - 0.75) < 1e-9:
                conv.append((math.log1p(y[i] / 4) + math.log1p(y[i] / 2)) / 0.75)
            else:
                conv.append(m * math.log1p(y[i] / m))
        d_conv = _stats([(knot[i] - conv[i]) * BP for i in range(len(weeks))])
        # (2) 엔진 vs 시트: 마디·주간
        d_knot = [(sc[w] - knot[i]) * BP for i, w in enumerate(weeks)]
        d_week = _stats([(sc[w] - week[w]) * BP for w in range(1, 521) if isinstance(week[w], (int, float))]) if week is not None else {"n": 0, "max": None, "mean": None}
        out[c] = {"sheet_knot_is_ytm_conversion_bp": d_conv, "engine_minus_sheet_knot_bp": {f"{w / 52:g}Y": round(d, 3) for w, d in zip(weeks, d_knot)}, "engine_minus_sheet_weekly_bp": d_week}
        print(f"  {c}: 시트 마디 현물 − YTM 변환 {_fmt(d_conv, 'bp')} → 시트는 부트스트랩 아님(YTM 을 현물로 간주)")
        print(f"     엔진 − 시트 마디(bp): " + " ".join(f"{k}:{v:+.2f}" for k, v in out[c]["engine_minus_sheet_knot_bp"].items()))
        print(f"     엔진 − 시트 주간 격자 {_fmt(d_week, 'bp')}")
    return out


# ----------------------------------------------------------------------------- D. 대교 'Par검증' 예제(템플릿) — 어떤 관례가 H·I = 1 을 만드는가
def part_D():
    from openpyxl import load_workbook
    wb = load_workbook(DAEKYO, read_only=True, data_only=True)
    rows = list(wb["Par검증"].iter_rows(min_row=5, max_row=8, max_col=8, values_only=True))
    wb.close()
    D30 = [r[2] for r in rows]; DACT = [r[3] for r in rows]; Y = [r[4] for r in rows]; S = [r[5] for r in rows]
    m = 4

    def h_nom(k, D):
        y = Y[k]; return sum((y / m) / (1 + y / m) ** (m * D[j]) for j in range(k + 1)) + 1 / (1 + y / m) ** (m * D[k])

    def i_ann(k, D):
        y = Y[k]; return sum((y / m) / (1 + S[j]) ** D[j] for j in range(k + 1)) + 1 / (1 + S[k]) ** D[k]

    def i_nom(k, D):
        y = Y[k]; return sum((y / m) / (1 + S[j] / m) ** (m * D[j]) for j in range(k + 1)) + 1 / (1 + S[k] / m) ** (m * D[k])

    def i_cont(k, D):
        y = Y[k]; return sum((y / m) * math.exp(-S[j] * D[j]) for j in range(k + 1)) + math.exp(-S[k] * D[k])

    out = {}
    print("\n[D] 대교 'Par검증' 예제(템플릿, 평가일 2020-12-31, 분기 4행): 각 관례로 계산한 H(YTM 항등식)·I(Spot 재가격)의 max|1−x|")
    for name, f in (("H_ytm_nominal_q", h_nom), ("I_spot_annual_eff(엔진 PAR_CHECK)", i_ann), ("I_spot_nominal_q", i_nom), ("I_spot_continuous", i_cont)):
        for dn, D in (("30/360", D30), ("Act/365F", DACT)):
            v = max(abs(1 - f(k, D)) for k in range(4))
            out[f"{name}|{dn}"] = v
            print(f"  {name:36} {dn:9} max|1−x| = {v:.2e}")
    return out


def part_C():
    return {"C1_YP_boot_strapping": part_C1(), "C2_KBI_BOOT": part_C2(), "C3_KBI_appraiser_sheet": part_C3()}


def main(argv):
    which = (argv[1] if len(argv) > 1 else "all").upper()
    out = {}
    if which in ("A", "ALL"):
        out["A"] = part_A(max_rd_rows=int(argv[2]) if len(argv) > 2 else None)
    if which in ("B", "ALL"):
        out["B"] = part_B()
    if which in ("C", "ALL"):
        out["C"] = part_C()
    if which in ("C1", "C2", "C3"):
        out[which] = globals()["part_" + which]()
    if which in ("D", "ALL"):
        out["D"] = part_D()
    os.makedirs("exports", exist_ok=True)
    with open(os.path.join("exports", "crosscheck_ref.json"), "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1, default=str)
    print("\nJSON → exports/crosscheck_ref.json")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))

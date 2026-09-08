# -*- coding: utf-8 -*-
"""
verify_reviewer_sheet.py — 개발 검증 도구(Excel 필요): 검토자 원본 통합문서의 입력(테너 주수·YTM·고시일)만으로 io/reviewer_sheet 가 만든
수식 시트를 실제 Excel(COM)로 재계산하고, 원본의 값복사 셀(전 열)과 셀 단위로 대조한다. 파이썬 모형(evaluate)도 같이 대조한다.
사용: python cb_valuation/step1_curve/reference/verify_reviewer_sheet.py <원본 xlsx> [출력 xlsx]
출력: 행별 max|차이|·정확 일치 수 — 원본은 값복사(해찾기 잔차 포함)라 DF 행 ~4e-15, 10000 배율 행 ~4e-11 수준 차이는 정상.
"""
from __future__ import annotations
import os, sys, tempfile
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
from cb_valuation.step1_curve.io import reviewer_sheet as RS  # noqa: E402
from cb_valuation.step1_curve.reference.excel_recalc import recalc  # noqa: E402


def col_index(s: str) -> int:
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def params_from_original(ws_vals, kind: str) -> RS.Params:
    steps = [int(ws_vals.cell(4, c).value) for c in range(3, 15)]
    ytm = [float(ws_vals.cell(6, c).value) for c in range(3, 15)]
    years = [w / 52 for w in steps]
    tenors = [{"label": RS._tenor_label(y), "years": y, "step": w, "ytm": v} for y, w, v in zip(years, steps, ytm)]
    cd = ws_vals.cell(5, 2).value
    return RS.Params(kind, 52, 4, 520, tenors, cd if isinstance(cd, datetime) else None, header="verify")


def main(argv):
    from openpyxl import load_workbook, Workbook
    src = argv[1]
    out = argv[2] if len(argv) > 2 else os.path.join(tempfile.gettempdir(), "reviewer_sheet_verify.xlsx")
    wbv = load_workbook(src, data_only=True)
    wb = Workbook(); wb.remove(wb.active)
    params = {}
    for kind, name in RS.SHEET_NAMES.items():
        p = params_from_original(wbv[name], kind); params[kind] = p
        RS.write_sheet(wb.create_sheet(name), wb, p)
    wb.save(out)
    recalc(out)
    wbr = load_workbook(out, data_only=True)
    wbf = load_workbook(out)
    worst = 0.0
    for kind, name in RS.SHEET_NAMES.items():
        p = params[kind]; V = RS.evaluate(p)
        wso, wsr, wsf = wbv[name], wbr[name], wbf[name]
        print(f"===== {name}")
        for r in range(1, p.rows["last"] + 1):
            diffs, exact, n, pydiff, errs = 0.0, 0, 0, 0.0, []
            for c in range(3, p.N + 4):
                o, v = wso.cell(r, c).value, wsr.cell(r, c).value
                if isinstance(o, bool) or isinstance(v, bool):
                    if o != v:
                        errs.append(f"{RS.col(c)}{r}: 원본 {o} 재계산 {v}")
                    continue
                if not isinstance(o, (int, float)) or not isinstance(v, (int, float)):
                    if isinstance(o, (int, float)) and not isinstance(v, (int, float)):
                        errs.append(f"{RS.col(c)}{r}: 원본 {o!r} 재계산 {v!r} ({wsf.cell(r, c).value!r})")
                    continue
                n += 1
                d = abs(o - v); diffs = max(diffs, d); exact += (o == v)
                pv = V.get(r)
                if c == p.N + 3:
                    pyv = V.get((r, "X"))
                else:
                    pyv = pv[c - 3] if pv is not None and c - 3 < len(pv) else None
                if isinstance(pyv, (int, float)) and not isinstance(pyv, bool):
                    pydiff = max(pydiff, abs(pyv - v))
            if n:
                worst = max(worst, diffs)
                print(f"r{r:>2} {str(wso.cell(r, 2).value)[:30]!s:32} n={n:>3} 원본-Excel max={diffs:.3e} exact={exact}/{n} | 모형-Excel max={pydiff:.3e}")
            for e in errs[:5]:
                print("   ", e)
    print("written", out, "| overall max |원본-Excel| =", f"{worst:.3e}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))

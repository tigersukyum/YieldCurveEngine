# -*- coding: utf-8 -*-
"""
verify_formula_sheet.py — 개발 검증 도구(Excel 필요): 프로필(LINEAR/PCHIP/REVIEWER_2024 …)로 fixture A 를 끝까지 실행해 export_xlsx 로
Rf_dc/Rd_dc **수식 시트**를 만들고, 실제 Excel(COM)로 재계산한 값을 파이썬 시트 모형(reviewer_sheet.evaluate)과 셀 단위로 대조한다.
PCHIP 기울기(행 8)·보조행·3차 Hermite(행 11) 수식이 참조 구현(reference/interp_ref)과 같은 값을 내는지 확인하는 용도.
사용: python cb_valuation/step1_curve/reference/verify_formula_sheet.py PROFILE [step=weekly] [horizon=10]
"""
from __future__ import annotations
import os, shutil, sys, tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
from cb_valuation.step1_curve.io import reviewer_sheet as RS  # noqa: E402
from cb_valuation.step1_curve.io.evidence_writer import export_xlsx  # noqa: E402
from cb_valuation.step1_curve.reference.excel_recalc import recalc  # noqa: E402
from cb_valuation.step1_curve.tests.test_app_pipeline import drive_to_done  # noqa: E402


def main(argv):
    from openpyxl import load_workbook
    profile = argv[1] if len(argv) > 1 else "PCHIP"
    step = argv[2] if len(argv) > 2 else "weekly"
    horizon = float(argv[3]) if len(argv) > 3 else 10.0
    base = tempfile.mkdtemp(prefix="cb_vfs_")
    try:
        s, C = drive_to_done(base, profile, step=step, horizon_years=horizon)
        print("profile", profile, "status", s.run.status, s.result.fail_reason, "| RF_FREQ/RD_FREQ", C.RF_FREQ, C.RD_FREQ, "| INTERP_METHOD_PRE", C.INTERP_METHOD_PRE)
        path = export_xlsx(s, C, base)
        recalc(path)
        wb = load_workbook(path, data_only=True)
        worst = 0.0
        for kind, name in RS.SHEET_NAMES.items():
            p = RS.params_from_state(s, C, kind); V = RS.evaluate(p); ws = wb[name]
            print(f"===== {name} (m={p.m}, spq={p.spq}, nT={p.nT}, method={p.method})")
            rows = sorted(r for r in V if isinstance(r, int))
            for r in rows:
                vals = V[r]; n = 0; mx = 0.0; errs = []
                for j, pyv in enumerate(vals):
                    v = ws.cell(r, j + 3).value
                    if isinstance(pyv, bool) or isinstance(v, bool):
                        if pyv != v:
                            errs.append(f"{RS.col(j + 3)}{r}: 모형 {pyv} Excel {v}")
                        continue
                    if not isinstance(pyv, (int, float)) or not isinstance(v, (int, float)):
                        if isinstance(pyv, (int, float)) and not isinstance(v, (int, float)):
                            errs.append(f"{RS.col(j + 3)}{r}: Excel 값 아님 {v!r}")
                        continue
                    n += 1; d = abs(pyv - v); mx = max(mx, d)
                x = V.get((r, "X"))
                if isinstance(x, (int, float)) and not isinstance(x, bool):
                    v = ws.cell(r, p.N + 3).value
                    if isinstance(v, (int, float)):
                        mx = max(mx, abs(x - v)); n += 1
                if n:
                    worst = max(worst, mx)
                    print(f"r{r:>2} {str(ws.cell(r, 2).value)[:30]!s:32} n={n:>3} 모형-Excel max={mx:.3e}")
                for e in errs[:5]:
                    print("   ", e)
        pc = wb["PAR_CHECK"]
        for r0, name in ((2, "국고채"), (None, "회사채")):
            if r0 is None:
                r0 = next(r for r in range(3, pc.max_row + 1) if str(pc.cell(r, 2).value or "").startswith("Par 검증 - 회사채"))
            rows = [r for r in range(r0 + 3, pc.max_row + 1) if isinstance(pc.cell(r, 6).value, (int, float)) and pc.cell(r, 6).value != 0 and not (isinstance(pc.cell(r, 3).value, str))]
            if not rows:
                print(f"PAR_CHECK {name}: 값 없음 (r0={r0})"); continue
            hs = [abs(pc.cell(r, 8).value - 1) for r in rows if isinstance(pc.cell(r, 8).value, (int, float))]
            is_ = [abs(pc.cell(r, 9).value - 1) for r in rows if isinstance(pc.cell(r, 9).value, (int, float))]
            zero = [r for r in range(r0 + 3, pc.max_row + 1) if pc.cell(r, 6).value == 0 and isinstance(pc.cell(r, 3).value, (int, float))]
            print(f"PAR_CHECK {name}: 행 {rows[0]}~{rows[-1]} ({len(rows)}개 값 있음, {len(zero)}개 산출 기간 밖) | max|H-1|={max(hs):.2e} max|I-1|={max(is_):.2e} | D 첫/끝 {pc.cell(rows[0], 4).value}/{pc.cell(rows[-1], 4).value}")
        print("workbook", path, "| overall max |모형-Excel| =", f"{worst:.3e}")
        keep = os.path.join(os.getcwd(), "exports", f"verify_{profile}_{step}.xlsx")
        os.makedirs(os.path.dirname(keep), exist_ok=True); shutil.copy(path, keep); print("copied to", keep)
    finally:
        shutil.rmtree(base, ignore_errors=True)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
